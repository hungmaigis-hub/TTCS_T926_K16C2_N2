import io
import pytest
import sys
from pathlib import Path
import openpyxl

# Đảm bảo import được module main từ thư mục backend
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from main import app
from database.session import SessionLocal
from database.models import HoSoThucTap, DanhGia

client = TestClient(app)

# ==============================================================================
# BỘ KIỂM THỬ CHO API: GET /api/v1/reports/export (Xuất báo cáo đa định dạng Excel / PDF)
# ==============================================================================

@pytest.fixture(autouse=True)
def setup_eval_data():
    """
    Đảm bảo dữ liệu mẫu hồ sơ và đánh giá luôn đầy đủ và nhất quán
    để việc xuất file có đủ dữ liệu test.
    """
    try:
        from tests.conftest import MYSQL_AVAILABLE, TestSessionLocal
        Session = SessionLocal if MYSQL_AVAILABLE else TestSessionLocal
    except Exception:
        Session = SessionLocal

    db = Session()
    try:
        hs1 = db.query(HoSoThucTap).filter(HoSoThucTap.ma_ho_so == 1).first()
        if hs1:
            hs1.ma_truong = 1
            hs1.chuyen_nganh = "Công nghệ thông tin"

        hs2 = db.query(HoSoThucTap).filter(HoSoThucTap.ma_ho_so == 2).first()
        if hs2:
            hs2.ma_truong = 2
            hs2.chuyen_nganh = "Khoa học máy tính"

        # Đảm bảo có ít nhất 1 bản ghi đánh giá
        dg = db.query(DanhGia).first()
        if not dg:
            new_dg = DanhGia(
                ma_ho_so=1,
                loai_danh_gia="CuoiKy",
                diem_ky_nang=9.0,
                diem_thai_do=9.5,
                diem_trung_binh=9.25,
                xep_loai="XuatSac",
                nhan_xet_chi_tiet="Thực tập sinh xuất sắc",
                de_xuat_tuyen_chinh_thuc=True,
            )
            db.add(new_dg)
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def test_reports_export_excel_success():
    """Kiểm tra xuất báo cáo định dạng Excel (.xlsx) thành công"""
    response = client.get("/api/v1/reports/export?format=excel")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]

    # Đọc nội dung workbook từ bytes trả về
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Tong_Hop_KPI" in wb.sheetnames
    assert "Chi_Tiet_Danh_Gia" in wb.sheetnames

    # Kiểm tra sheet tổng hợp KPI
    ws_kpi = wb["Tong_Hop_KPI"]
    assert ws_kpi["A1"].value == "BÁO CÁO TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ THỰC TẬP SINH"

    # Kiểm tra sheet chi tiết
    ws_detail = wb["Chi_Tiet_Danh_Gia"]
    assert ws_detail["A1"].value == "STT"
    assert ws_detail["D1"].value == "Họ và tên"
    assert ws_detail.max_row >= 2


def test_reports_export_default_format_is_excel():
    """Kiểm tra mặc định khi không truyền format thì xuất file Excel (.xlsx)"""
    response = client.get("/api/v1/reports/export")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Tong_Hop_KPI" in wb.sheetnames


def test_reports_export_xlsx_alias():
    """Kiểm tra format=xlsx hoạt động như format=excel"""
    response = client.get("/api/v1/reports/export?format=xlsx")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert len(response.content) > 1000


def test_reports_export_pdf_success():
    """Kiểm tra xuất báo cáo định dạng PDF (.pdf) bằng ReportLab thành công"""
    response = client.get("/api/v1/reports/export?format=pdf")
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    assert ".pdf" in response.headers["content-disposition"]

    # Kiểm tra magic bytes đặc trưng của file PDF
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 1000


def test_reports_export_excel_with_filters():
    """Kiểm tra xuất báo cáo Excel khi kết hợp bộ lọc (ma_truong, loai_danh_gia)"""
    response = client.get("/api/v1/reports/export?format=excel&ma_truong=1&loai_danh_gia=CuoiKy")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws_detail = wb["Chi_Tiet_Danh_Gia"]
    assert ws_detail.max_row >= 2


def test_reports_export_pdf_with_filters():
    """Kiểm tra xuất báo cáo PDF khi kết hợp bộ lọc (de_xuat_tuyen_chinh_thuc)"""
    response = client.get("/api/v1/reports/export?format=pdf&de_xuat_tuyen_chinh_thuc=true")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"


def test_reports_export_invalid_format():
    """Kiểm tra truyền format không hợp lệ (ví dụ format=word) trả về HTTP 400"""
    response = client.get("/api/v1/reports/export?format=word")
    assert response.status_code == 400
    assert "Định dạng xuất file không được hỗ trợ" in response.json()["detail"]


def test_reports_export_invalid_format_csv():
    """Kiểm tra truyền format không hợp lệ (format=csv) trả về HTTP 400"""
    response = client.get("/api/v1/reports/export?format=csv")
    assert response.status_code == 400
    assert "Định dạng xuất file không được hỗ trợ" in response.json()["detail"]


def test_reports_export_empty_data_excel():
    """Kiểm tra khi lọc dữ liệu không có kết quả, file Excel vẫn xuất thành công không lỗi"""
    response = client.get("/api/v1/reports/export?format=excel&tu_khoa=KhongTheNaoTimThayTenNay123456")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws_detail = wb["Chi_Tiet_Danh_Gia"]
    assert ws_detail["A2"].value == "Không có dữ liệu đánh giá phù hợp"


def test_reports_export_empty_data_pdf():
    """Kiểm tra khi lọc dữ liệu không có kết quả, file PDF vẫn xuất an toàn không crash"""
    response = client.get("/api/v1/reports/export?format=pdf&tu_khoa=KhongTheNaoTimThayTenNay123456")
    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"
