import io
import sys
from pathlib import Path
from datetime import datetime
import pytest
import openpyxl
import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from main import app
from database.models import ChuongTrinhThucTap, NguoiDung, HoSoThucTap
from database.session import SessionLocal

client = TestClient(app)


def get_db():
    try:
        from tests.conftest import TestSessionLocal
        return TestSessionLocal()
    except Exception:
        return SessionLocal()


def create_excel_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test"
    headers = ["STT", "Mã SV", "Họ và tên", "Email", "Số điện thoại", "Lớp", "Chuyên ngành", "Trường"]
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def create_docx_bytes(rows):
    doc = docx.Document()
    headers = ["STT", "Mã SV", "Họ và tên", "Email", "Số điện thoại", "Lớp", "Chuyên ngành", "Trường"]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(r):
            row_cells[i].text = str(val)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_download_template_excel():
    response = client.get("/api/v1/interns/template-excel")
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers.get("content-type", "")
    assert len(response.content) > 0


def test_download_template_docx():
    response = client.get("/api/v1/interns/template-docx")
    assert response.status_code == 200
    assert "wordprocessingml.document" in response.headers.get("content-type", "")
    assert len(response.content) > 0


def test_import_preview_mode():
    db = get_db()
    prog = db.query(ChuongTrinhThucTap).first()
    assert prog is not None

    rows = [
        [1, "TEST001", "Nguyễn Văn Hợp Lệ", f"hop_le_{datetime.now().timestamp()}@ictu.edu.vn", "0987654321", "K19", "CNTT", "ICTU"],
        [2, "TEST002", "Trần Lỗi Email", "email_sai_dinh_dang", "0987654322", "K19", "CNTT", "ICTU"],
        [3, "TEST003", "", f"thieu_ten_{datetime.now().timestamp()}@ictu.edu.vn", "0987654323", "K19", "CNTT", "ICTU"]
    ]
    file_bytes = create_excel_bytes(rows)

    response = client.post(
        "/api/v1/interns/import",
        data={
            "ma_chuong_trinh": prog.ma_chuong_trinh,
            "che_do": "kiem_tra_truoc"
        },
        files={"file": ("test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["data"]["che_do"] == "kiem_tra_truoc"
    assert data["data"]["da_luu"] is False
    assert data["data"]["thanh_cong"] == 1
    assert data["data"]["that_bai"] == 2
    assert len(data["data"]["chi_tiet_loi"]) == 2


def test_import_strict_mode_with_errors():
    db = get_db()
    prog = db.query(ChuongTrinhThucTap).first()
    assert prog is not None

    rows = [
        [1, "STRICT01", "Sinh Viên Một", f"strict1_{datetime.now().timestamp()}@ictu.edu.vn", "0987654321", "K19", "CNTT", "ICTU"],
        [2, "STRICT02", "Sinh Viên Sai", "email_khong_hop_le", "0987654322", "K19", "CNTT", "ICTU"]
    ]
    file_bytes = create_excel_bytes(rows)

    response = client.post(
        "/api/v1/interns/import",
        data={
            "ma_chuong_trinh": prog.ma_chuong_trinh,
            "che_do": "dung_khi_co_loi"
        },
        files={"file": ("test_strict.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 400
    data = response.json()
    assert data["data"]["da_luu"] is False
    assert data["data"]["that_bai"] >= 1


def test_import_skip_errors_mode_excel():
    db = get_db()
    prog = db.query(ChuongTrinhThucTap).first()
    assert prog is not None

    unique_email = f"pass_{datetime.now().timestamp()}@ictu.edu.vn"
    rows = [
        [1, f"SV{int(datetime.now().timestamp())}", "Sinh Viên Chuẩn", unique_email, "0987111222", "K19-KTPM", "Kỹ thuật phần mềm", "ICTU"],
        [2, "SV_FAIL", "Sinh Viên Lỗi Email", "bad-email", "0987333444", "K19", "CNTT", "ICTU"]
    ]
    file_bytes = create_excel_bytes(rows)

    response = client.post(
        "/api/v1/interns/import",
        data={
            "ma_chuong_trinh": prog.ma_chuong_trinh,
            "che_do": "bo_qua_loi"
        },
        files={"file": ("test_skip.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["data"]["da_luu"] is True
    assert data["data"]["thanh_cong"] >= 1
    assert data["data"]["that_bai"] >= 1

    # Kiểm tra đã lưu vào CSDL
    user = db.query(NguoiDung).filter(NguoiDung.email == unique_email).first()
    assert user is not None
    hs = db.query(HoSoThucTap).filter(HoSoThucTap.ma_nguoi_dung == user.ma_nguoi_dung).first()
    assert hs is not None


def test_import_docx_format():
    db = get_db()
    prog = db.query(ChuongTrinhThucTap).first()
    assert prog is not None

    unique_email = f"docx_{datetime.now().timestamp()}@ictu.edu.vn"
    rows = [
        ["1", f"DOCX{int(datetime.now().timestamp())}", "Sinh Viên Word", unique_email, "0977888999", "K20", "Khoa học máy tính", "ICTU"]
    ]
    file_bytes = create_docx_bytes(rows)

    response = client.post(
        "/api/v1/interns/import",
        data={
            "ma_chuong_trinh": prog.ma_chuong_trinh,
            "che_do": "bo_qua_loi"
        },
        files={"file": ("test.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["data"]["thanh_cong"] == 1


def test_export_errors_to_excel():
    chi_tiet_loi = [
        {"dong": 2, "ma_sinh_vien": "DTC01", "ho_ten": "Nguyễn Văn A", "email": "sai_email", "so_dien_thoai": "0987654321", "ly_do": "Email sai định dạng"},
        {"dong": 5, "ma_sinh_vien": "DTC02", "ho_ten": "--", "email": "a@ictu.edu.vn", "so_dien_thoai": "0987654322", "ly_do": "Thiếu họ tên"}
    ]
    response = client.post(
        "/api/v1/interns/export-errors-excel",
        json={"chi_tiet_loi": chi_tiet_loi}
    )
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers.get("content-type", "")
    assert len(response.content) > 0


def test_import_pdf_format():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    db = get_db()
    prog = db.query(ChuongTrinhThucTap).first()
    assert prog is not None

    unique_email = f"pdf_{datetime.now().timestamp()}@ictu.edu.vn"
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(40, 750, "STT | Mã SV | Họ và tên | Email | Số điện thoại | Lớp | Chuyên ngành | Trường")
    c.drawString(40, 720, f"1 | PDF001 | Sinh Viên PDF | {unique_email} | 0988777666 | K19 | CNTT | ICTU")
    c.save()
    buf.seek(0)
    pdf_bytes = buf.getvalue()

    response = client.post(
        "/api/v1/interns/import",
        data={
            "ma_chuong_trinh": prog.ma_chuong_trinh,
            "che_do": "bo_qua_loi"
        },
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["data"]["thanh_cong"] >= 1


if __name__ == "__main__":
    test_download_template_excel()
    test_download_template_docx()
    test_import_preview_mode()
    test_import_strict_mode_with_errors()
    test_import_skip_errors_mode_excel()
    test_import_docx_format()
    test_export_errors_to_excel()
    test_import_pdf_format()
    print("ALL 8 IMPORT TESTS PASSED SUCCESSFULLY!")


