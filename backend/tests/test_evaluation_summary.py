import io
import pytest
import sys
from pathlib import Path
import openpyxl

# Đảm bảo import được module main từ thư mục backend
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from main import app
from database.session import SessionLocal
from database.models import HoSoThucTap, DanhGia

client = TestClient(app)

# ==============================================================================
# BỘ KIỂM THỬ CHO API: GET /api/v1/evaluations/summary (Tổng hợp dữ liệu đánh giá & Xuất Excel/PDF)
# ==============================================================================

@pytest.fixture(autouse=True)
def setup_eval_data():
    """
    Đảm bảo dữ liệu mẫu hồ sơ và đánh giá luôn đầy đủ và nhất quán
    (phòng trường hợp các test case khác như PUT làm thay đổi ma_truong hoặc chuyen_nganh)
    """
    try:
        from tests.conftest import MYSQL_AVAILABLE, TestSessionLocal
        Session = SessionLocal if MYSQL_AVAILABLE else TestSessionLocal
    except Exception:
        Session = SessionLocal

    db = Session()
    try:
        hs1 = db.query(HoSoThucTap).filter(HoSoThucTap.ma_ho_so == 1).first()
        if hs1:
            hs1.ma_truong = 1
            hs1.chuyen_nganh = "Công nghệ thông tin"

        hs2 = db.query(HoSoThucTap).filter(HoSoThucTap.ma_ho_so == 2).first()
        if hs2:
            hs2.ma_truong = 2
            hs2.chuyen_nganh = "Khoa học máy tính"

        # Đảm bảo 4 bản ghi đánh giá luôn tồn tại
        if db.query(DanhGia).count() < 4:
            db.query(DanhGia).delete()
            dg1 = DanhGia(ma_danh_gia=1, ma_ho_so=1, ma_nguoi_danh_gia=3, loai_danh_gia="GiuaKy", diem_ky_nang=8.5, diem_thai_do=9.0, nhan_xet_chi_tiet="Tiếp thu nhanh, hoàn thành tốt nhiệm vụ", de_xuat_tuyen_chinh_thuc=False)
            dg2 = DanhGia(ma_danh_gia=2, ma_ho_so=1, ma_nguoi_danh_gia=3, loai_danh_gia="CuoiKy", diem_ky_nang=9.0, diem_thai_do=9.5, nhan_xet_chi_tiet="Kỹ năng chuyên môn xuất sắc, trách nhiệm cao", de_xuat_tuyen_chinh_thuc=True)
            dg3 = DanhGia(ma_danh_gia=3, ma_ho_so=2, ma_nguoi_danh_gia=3, loai_danh_gia="GiuaKy", diem_ky_nang=7.5, diem_thai_do=8.0, nhan_xet_chi_tiet="Thực hiện công việc đúng tiến độ", de_xuat_tuyen_chinh_thuc=False)
            dg4 = DanhGia(ma_danh_gia=4, ma_ho_so=2, ma_nguoi_danh_gia=3, loai_danh_gia="CuoiKy", diem_ky_nang=8.0, diem_thai_do=8.5, nhan_xet_chi_tiet="Tiến bộ rõ rệt, đáp ứng tốt yêu cầu", de_xuat_tuyen_chinh_thuc=False)
            db.add_all([dg1, dg2, dg3, dg4])

        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()
    yield


def test_evaluation_summary_default_json():
    """Kiểm tra gọi API lấy dữ liệu JSON mặc định -> 200 OK kèm đầy đủ cấu trúc"""
    response = client.get("/api/v1/evaluations/summary")
    assert response.status_code == 200

    data = response.json()
    assert data["status_code"] == 200
    assert data["message"] == "Lấy dữ liệu tổng hợp đánh giá thành công"
    assert "data" in data

    res_body = data["data"]
    assert "summary" in res_body
    assert "items" in res_body
    assert "pagination" in res_body

    summary = res_body["summary"]
    assert summary["tong_so_danh_gia"] >= 4
    assert summary["tong_so_ho_so"] >= 2
    assert "diem_ky_nang_tb" in summary
    assert "diem_thai_do_tb" in summary
    assert "diem_tong_ket_tb" in summary
    assert "so_luong_de_xuat_tuyen_dung" in summary
    assert "ty_le_de_xuat_tuyen_dung" in summary
    assert "phan_bo_xep_loai" in summary
    assert "thong_ke_theo_truong" in summary


def test_evaluation_summary_kpi_calculation_accuracy():
    """Kiểm tra độ chính xác của các chỉ số KPI tính toán"""
    response = client.get("/api/v1/evaluations/summary")
    assert response.status_code == 200

    summary = response.json()["data"]["summary"]
    items = response.json()["data"]["items"]

    total = summary["tong_so_danh_gia"]
    assert total > 0

    # Kiểm tra điểm kỹ năng TB
    calc_ky_nang = round(sum(it["diem_ky_nang"] for it in items) / len(items), 2)
    assert abs(summary["diem_ky_nang_tb"] - calc_ky_nang) < 0.05

    # Kiểm tra điểm thái độ TB
    calc_thai_do = round(sum(it["diem_thai_do"] for it in items) / len(items), 2)
    assert abs(summary["diem_thai_do_tb"] - calc_thai_do) < 0.05

    # Kiểm tra tỷ lệ đề xuất tuyển dụng
    de_xuat_count = sum(1 for it in items if it["de_xuat_tuyen_chinh_thuc"])
    assert summary["so_luong_de_xuat_tuyen_dung"] == de_xuat_count
    assert summary["ty_le_de_xuat_tuyen_dung"] == round((de_xuat_count / total) * 100.0, 2)


def test_evaluation_summary_items_fields():
    """Kiểm tra chi tiết từng bản ghi trong danh sách items có đủ thông tin các bảng liên kết"""
    response = client.get("/api/v1/evaluations/summary")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    item = items[0]

    expected_fields = [
        "ma_danh_gia", "ma_ho_so", "ho_ten", "email", "so_dien_thoai",
        "chuyen_nganh", "ma_truong", "ten_truong", "trang_thai_thuc_tap",
        "loai_danh_gia", "diem_ky_nang", "diem_thai_do", "diem_trung_binh",
        "xep_loai", "nhan_xet_chi_tiet", "de_xuat_tuyen_chinh_thuc", "nguoi_danh_gia"
    ]
    for field in expected_fields:
        assert field in item, f"Trường {field} thiếu trong item"


def test_evaluation_summary_filter_by_ma_truong():
    """Kiểm tra lọc theo mã trường đại học hợp lệ"""
    response = client.get("/api/v1/evaluations/summary?ma_truong=1")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert it["ma_truong"] == 1


def test_evaluation_summary_filter_by_ma_truong_nonexistent():
    """Kiểm tra lọc theo mã trường không tồn tại -> trả về rỗng nhưng vẫn 200 OK"""
    response = client.get("/api/v1/evaluations/summary?ma_truong=999999")
    assert response.status_code == 200

    res_body = response.json()["data"]
    assert res_body["summary"]["tong_so_danh_gia"] == 0
    assert res_body["summary"]["tong_so_ho_so"] == 0
    assert len(res_body["items"]) == 0


def test_evaluation_summary_filter_by_loai_danh_gia_giuaky():
    """Kiểm tra lọc theo loại đánh giá GiuaKy"""
    response = client.get("/api/v1/evaluations/summary?loai_danh_gia=GiuaKy")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert "GiuaKy" in it["loai_danh_gia"]


def test_evaluation_summary_filter_by_loai_danh_gia_cuoiky():
    """Kiểm tra lọc theo loại đánh giá CuoiKy"""
    response = client.get("/api/v1/evaluations/summary?loai_danh_gia=CuoiKy")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert "CuoiKy" in it["loai_danh_gia"]


def test_evaluation_summary_filter_by_trang_thai_thuc_tap():
    """Kiểm tra lọc theo trạng thái thực tập của sinh viên"""
    response = client.get("/api/v1/evaluations/summary?trang_thai_thuc_tap=DangThucTap")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    for it in items:
        assert "DangThucTap" in it["trang_thai_thuc_tap"]


def test_evaluation_summary_filter_by_de_xuat_true():
    """Kiểm tra lọc các bản ghi có đề xuất tuyển dụng chính thức = true"""
    response = client.get("/api/v1/evaluations/summary?de_xuat_tuyen_chinh_thuc=true")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert it["de_xuat_tuyen_chinh_thuc"] is True


def test_evaluation_summary_filter_by_de_xuat_false():
    """Kiểm tra lọc các bản ghi đề xuất tuyển dụng = false"""
    response = client.get("/api/v1/evaluations/summary?de_xuat_tuyen_chinh_thuc=false")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert it["de_xuat_tuyen_chinh_thuc"] is False


def test_evaluation_summary_search_by_tu_khoa_name():
    """Kiểm tra tìm kiếm từ khóa theo tên sinh viên"""
    response = client.get("/api/v1/evaluations/summary?tu_khoa=Nguy%E1%BB%85n")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert "Nguyễn" in it["ho_ten"] or "Nguyễn" in it["ten_truong"]


def test_evaluation_summary_search_by_tu_khoa_email():
    """Kiểm tra tìm kiếm từ khóa theo email sinh viên"""
    response = client.get("/api/v1/evaluations/summary?tu_khoa=example.com")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0
    for it in items:
        assert "example.com" in it["email"]


def test_evaluation_summary_search_by_tu_khoa_major():
    """Kiểm tra tìm kiếm từ khóa theo chuyên ngành đào tạo"""
    response = client.get("/api/v1/evaluations/summary?tu_khoa=th%C3%B4ng+tin")
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    assert len(items) > 0


def test_evaluation_summary_search_no_match():
    """Kiểm tra tìm kiếm từ khóa không có kết quả khớp"""
    response = client.get("/api/v1/evaluations/summary?tu_khoa=KHONG_TON_TAI_XYZ123")
    assert response.status_code == 200

    data = response.json()["data"]
    assert data["summary"]["tong_so_danh_gia"] == 0
    assert len(data["items"]) == 0


def test_evaluation_summary_combined_filters():
    """Kiểm tra kết hợp nhiều bộ lọc cùng lúc"""
    response = client.get(
        "/api/v1/evaluations/summary?ma_truong=1&loai_danh_gia=CuoiKy&de_xuat_tuyen_chinh_thuc=true"
    )
    assert response.status_code == 200

    items = response.json()["data"]["items"]
    for it in items:
        assert it["ma_truong"] == 1
        assert "CuoiKy" in it["loai_danh_gia"]
        assert it["de_xuat_tuyen_chinh_thuc"] is True


def test_evaluation_summary_pagination():
    """Kiểm tra phân trang dữ liệu khi lấy định dạng JSON"""
    # Lấy page 1 với page_size = 2
    res_p1 = client.get("/api/v1/evaluations/summary?page=1&page_size=2")
    assert res_p1.status_code == 200
    p1_data = res_p1.json()["data"]
    assert len(p1_data["items"]) == 2
    assert p1_data["pagination"]["page"] == 1
    assert p1_data["pagination"]["page_size"] == 2
    assert p1_data["pagination"]["total_pages"] >= 2

    # Lấy page 2 với page_size = 2
    res_p2 = client.get("/api/v1/evaluations/summary?page=2&page_size=2")
    assert res_p2.status_code == 200
    p2_data = res_p2.json()["data"]
    assert len(p2_data["items"]) == 2
    assert p2_data["pagination"]["page"] == 2

    # Đảm bảo 2 trang không trùng nhau
    p1_ids = [it["ma_danh_gia"] for it in p1_data["items"]]
    p2_ids = [it["ma_danh_gia"] for it in p2_data["items"]]
    assert set(p1_ids).isdisjoint(set(p2_ids))


def test_evaluation_summary_pagination_out_of_range():
    """Kiểm tra yêu cầu số trang vượt quá giới hạn -> trả về danh sách rỗng"""
    response = client.get("/api/v1/evaluations/summary?page=9999&page_size=50")
    assert response.status_code == 200
    assert len(response.json()["data"]["items"]) == 0


def test_evaluation_summary_invalid_page():
    """Kiểm tra page <= 0 trả về lỗi validation 422"""
    response = client.get("/api/v1/evaluations/summary?page=0")
    assert response.status_code == 422


def test_evaluation_summary_invalid_page_size():
    """Kiểm tra page_size vượt quá 100 trả về lỗi validation 422"""
    response = client.get("/api/v1/evaluations/summary?page_size=101")
    assert response.status_code == 422


def test_evaluation_summary_invalid_format():
    """Kiểm tra format không được hỗ trợ -> 400 Bad Request"""
    response = client.get("/api/v1/evaluations/summary?format=docx")
    assert response.status_code == 400
    assert "Định dạng xuất file không được hỗ trợ" in response.json()["detail"]


def test_evaluation_summary_export_excel_success():
    """Kiểm tra xuất file Excel thành công (.xlsx)"""
    response = client.get("/api/v1/evaluations/summary?format=excel")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]

    # Đọc nội dung workbook từ bytes response
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Tong_Hop_KPI" in wb.sheetnames
    assert "Chi_Tiet_Danh_Gia" in wb.sheetnames

    # Kiểm tra sheet KPI
    ws_kpi = wb["Tong_Hop_KPI"]
    assert ws_kpi["A1"].value == "BÁO CÁO TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ THỰC TẬP SINH"

    # Kiểm tra sheet chi tiết
    ws_detail = wb["Chi_Tiet_Danh_Gia"]
    assert ws_detail["A1"].value == "STT"
    assert ws_detail["D1"].value == "Họ và tên"
    assert ws_detail.max_row >= 5 # Ít nhất dòng tiêu đề + 4 bản ghi


def test_evaluation_summary_export_xlsx_alias():
    """Kiểm tra tham số format=xlsx hoạt động đồng nhất với format=excel"""
    response = client.get("/api/v1/evaluations/summary?format=xlsx")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert len(response.content) > 1000


def test_evaluation_summary_export_pdf_success():
    """Kiểm tra xuất file PDF thành công (.pdf)"""
    response = client.get("/api/v1/evaluations/summary?format=pdf")
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    assert ".pdf" in response.headers["content-disposition"]

    # Kiểm tra magic bytes của định dạng PDF
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 1000


def test_evaluation_summary_export_excel_with_filters():
    """Kiểm tra xuất file Excel khi kết hợp bộ lọc"""
    response = client.get("/api/v1/evaluations/summary?format=excel&ma_truong=1&loai_danh_gia=CuoiKy")
    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws_detail = wb["Chi_Tiet_Danh_Gia"]
    # Kiểm tra ít nhất 1 dòng dữ liệu + header
    assert ws_detail.max_row >= 2


def test_evaluation_summary_export_pdf_with_filters():
    """Kiểm tra xuất file PDF khi kết hợp bộ lọc"""
    response = client.get("/api/v1/evaluations/summary?format=pdf&de_xuat_tuyen_chinh_thuc=true")
    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"
