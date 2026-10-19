import io
import re
import csv
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import docx
except ImportError:
    docx = None

try:
    import pypdf
except ImportError:
    pypdf = None



EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$")
PHONE_REGEX = re.compile(r"^(0|\+84)[35789][0-9]{8}$")


def generate_intern_import_template() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mau_Nhap_Sinh_Vien"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Segoe UI", size=15, bold=True, color="1F497D")
    note_font = Font(name="Segoe UI", size=9, italic=True, color="595959")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    example_font = Font(name="Segoe UI", size=10, italic=True, color="333333")

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = "DANH SÁCH SINH VIÊN TIẾP NHẬN THỰC TẬP (MẪU IMPORT)"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    ws["A2"] = "Quy chuẩn: Các cột có dấu (*) là bắt buộc. Email và Số điện thoại phải đúng định dạng chuẩn. Không xóa hoặc sửa đổi dòng tiêu đề (Dòng 4)."
    ws["A2"].font = note_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    headers = [
        "STT",
        "Mã Sinh Viên (*)",
        "Họ Và Tên (*)",
        "Email (*)",
        "Số Điện Thoại (*)",
        "Lớp (*)",
        "Chuyên Ngành (*)",
        "Trường Đại Học"
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    sample_rows = [
        [1, "DTC2051060101", "Nguyễn Văn An", "an.nv@ictu.edu.vn", "0987654321", "K19-KTPM01", "Kỹ thuật phần mềm", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"],
        [2, "DTC2051060102", "Trần Thị Mai", "mai.tt@ictu.edu.vn", "0912345678", "K19-CNTT02", "Công nghệ thông tin", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"],
        [3, "DTC2051060103", "Lê Hoàng Long", "long.lh@ictu.edu.vn", "0978123456", "K20-KHMT01", "Khoa học máy tính", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"]
    ]

    for r_idx, row_data in enumerate(sample_rows, 5):
        ws.row_dimensions[r_idx].height = 22
        fill = alt_fill if r_idx % 2 == 1 else None
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = example_font
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if c_idx in (1, 2, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    col_widths = {1: 8, 2: 20, 3: 26, 4: 28, 5: 18, 6: 16, 7: 24, 8: 38}
    for col_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_intern_import_template_docx() -> io.BytesIO:
    if docx is None:
        raise ValueError("Thư viện python-docx chưa được cài đặt trong môi trường máy chủ")
    doc = docx.Document()

    p_title = doc.add_paragraph()
    p_title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("DANH SÁCH SINH VIÊN TIẾP NHẬN THỰC TẬP (MẪU IMPORT)")
    run_title.bold = True
    run_title.font.size = docx.shared.Pt(14)
    run_title.font.name = "Segoe UI"

    p_note = doc.add_paragraph()
    p_note.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    run_note = p_note.add_run("Quy chuẩn: Các cột có dấu (*) là bắt buộc. Không sửa đổi các cột ở hàng tiêu đề của bảng bên dưới.")
    run_note.italic = True
    run_note.font.size = docx.shared.Pt(9.5)
    run_note.font.name = "Segoe UI"

    headers = [
        "STT",
        "Mã Sinh Viên (*)",
        "Họ Và Tên (*)",
        "Email (*)",
        "Số Điện Thoại (*)",
        "Lớp (*)",
        "Chuyên Ngành (*)",
        "Trường Đại Học"
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = "Segoe UI"
                r.font.size = docx.shared.Pt(9.5)

    sample_rows = [
        ["1", "DTC2051060101", "Nguyễn Văn An", "an.nv@ictu.edu.vn", "0987654321", "K19-KTPM01", "Kỹ thuật phần mềm", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"],
        ["2", "DTC2051060102", "Trần Thị Mai", "mai.tt@ictu.edu.vn", "0912345678", "K19-CNTT02", "Công nghệ thông tin", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"],
        ["3", "DTC2051060103", "Lê Hoàng Long", "long.lh@ictu.edu.vn", "0978123456", "K20-KHMT01", "Khoa học máy tính", "Trường ĐH Công nghệ Thông tin & Truyền thông (ICTU)"]
    ]

    for row_data in sample_rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            for p in row_cells[i].paragraphs:
                for r in p.runs:
                    r.font.name = "Segoe UI"
                    r.font.size = docx.shared.Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def export_error_interns_to_excel(chi_tiet_loi: List[Dict[str, Any]]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh_Sach_Loi_Import"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Segoe UI", size=14, bold=True, color="C00000")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    reason_font = Font(name="Segoe UI", size=10, bold=True, color="C00000")

    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    alt_fill = PatternFill(start_color="FFF0F0", end_color="FFF0F0", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.merge_cells("A1:F1")
    ws["A1"] = f"DANH SÁCH SINH VIÊN BỊ LỖI KHI NHẬP FILE ({len(chi_tiet_loi)} DÒNG BỊ LỖI)"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"] = "Hướng dẫn: Sửa lại các trường thông tin bị sai sót theo cột 'Lý Do Lỗi' rồi lưu file và import lại vào hệ thống."
    ws["A2"].font = Font(name="Segoe UI", size=9.5, italic=True, color="595959")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    headers = ["Dòng Trong File", "Mã Sinh Viên", "Họ Và Tên", "Email", "Số Điện Thoại", "Lý Do Bị Lỗi"]
    ws.row_dimensions[4].height = 25
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for idx, item in enumerate(chi_tiet_loi, 1):
        r_idx = 4 + idx
        ws.row_dimensions[r_idx].height = 22
        fill = alt_fill if idx % 2 == 1 else None

        row_vals = [
            item.get("dong", idx),
            item.get("ma_sinh_vien") or "--",
            item.get("ho_ten") or "--",
            item.get("email") or "--",
            item.get("so_dien_thoai") or "--",
            item.get("ly_do") or "Không rõ nguyên nhân"
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = reason_font if c_idx == 6 else data_font
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if c_idx in (1, 2):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    col_widths = {1: 16, 2: 20, 3: 26, 4: 30, 5: 18, 6: 50}
    for col_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_interns_to_excel(interns: List[Dict[str, Any]], program_name: Optional[str] = None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh_Sach_Thuc_Tap_Sinh"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Segoe UI", size=15, bold=True, color="1F497D")
    sub_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    ws.merge_cells("A1:K1")
    ws["A1"] = "DANH SÁCH HỒ SƠ THỰC TẬP SINH DOANH NGHIỆP"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Tổng số hồ sơ: {len(interns)}" + (f" | Chương trình: {program_name}" if program_name else "")
    ws["A2"].font = sub_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    headers = [
        "STT",
        "Mã Hồ Sơ",
        "Mã Sinh Viên",
        "Họ Và Tên",
        "Email",
        "Số Điện Thoại",
        "Lớp",
        "Chuyên Ngành",
        "Chương Trình Thực Tập",
        "Mentor Hướng Dẫn",
        "Trạng Thái"
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for idx, item in enumerate(interns, 1):
        r_idx = 4 + idx
        ws.row_dimensions[r_idx].height = 20
        fill = alt_fill if idx % 2 == 1 else None

        row_vals = [
            idx,
            f"#{item.get('ma_ho_so')}",
            item.get("ma_sinh_vien") or (item.get("email", "").split("@")[0].upper() if item.get("email") else "--"),
            item.get("ho_ten") or "--",
            item.get("email") or "--",
            item.get("so_dien_thoai") or "--",
            item.get("lop") or "--",
            item.get("chuyen_nganh") or "--",
            item.get("ten_chuong_trinh") or "--",
            item.get("ten_mentor") or "Chưa phân công",
            item.get("trang_thai_thuc_tap") or item.get("trang_thai_xet_duyet") or "--"
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if c_idx in (1, 2, 3, 6, 7, 11):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    col_widths = {1: 6, 2: 12, 3: 18, 4: 25, 5: 28, 6: 16, 7: 15, 8: 22, 9: 30, 10: 24, 11: 18}
    for col_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _normalize_header(text: str) -> str:
    if not text:
        return ""
    t = str(text).strip().lower()
    t = re.sub(r"[\s\(\)\*\:\_]+", " ", t).strip()
    return t


def _detect_column_mapping(headers: List[str]) -> Dict[str, int]:
    mapping = {}
    has_stt = False
    for idx, raw in enumerate(headers):
        h = _normalize_header(raw)
        if any(k in h for k in ["stt", "số thứ tự", "no.", "no", "order", "index"]):
            has_stt = True
        elif any(k in h for k in ["mã số sinh viên", "mã sinh viên", "mã sv", "masv", "mssv", "mã_sv", "student id", "student_code", "mã"]):
            mapping["ma_sinh_vien"] = idx
        elif any(k in h for k in ["họ và tên", "họ tên", "họ & tên", "fullname", "full name", "tên sinh viên", "họ tên sv", "họ và tên sinh viên"]):
            mapping["ho_ten"] = idx
        elif any(k in h for k in ["email", "thư điện tử", "e-mail", "mail", "hòm thư"]):
            mapping["email"] = idx
        elif any(k in h for k in ["số điện thoại", "sđt", "sdt", "điện thoại", "phone", "mobile", "đtdđ", "di động", "số đt"]):
            mapping["so_dien_thoai"] = idx
        elif any(k in h for k in ["lớp", "class", "chi đoàn", "lớp sinh hoạt", "lớp quản lý", "lớp học"]):
            mapping["lop"] = idx
        elif any(k in h for k in ["chuyên ngành", "ngành học", "ngành", "khoa/ngành", "major"]):
            mapping["chuyen_nganh"] = idx
        elif any(k in h for k in ["trường đại học", "trường", "đại học", "university", "cơ sở đào tạo", "đơn vị đào tạo", "học viện"]):
            mapping["truong"] = idx

    if "ho_ten" not in mapping:
        for idx, raw in enumerate(headers):
            h = _normalize_header(raw)
            if "tên" in h and "trường" not in h:
                mapping["ho_ten"] = idx
                break

    if "ho_ten" not in mapping and len(headers) >= 3:
        offset = 1 if has_stt else 0
        mapping.setdefault("ma_sinh_vien", offset)
        mapping.setdefault("ho_ten", offset + 1)
        mapping.setdefault("email", offset + 2)
        mapping.setdefault("so_dien_thoai", offset + 3)
        mapping.setdefault("lop", offset + 4)
        mapping.setdefault("chuyen_nganh", offset + 5)
        mapping.setdefault("truong", offset + 6)

    return mapping


def parse_intern_file(file_bytes: bytes, filename: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    ext = filename.lower().split(".")[-1]
    raw_rows: List[Dict[str, Any]] = []
    file_errors: List[str] = []

    if ext in ("xlsx", "xls"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active

            header_row_idx = None
            header_vals = []
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
                norm_vals = [_normalize_header(v) for v in row_vals]
                non_empty = [v for v in norm_vals if v]
                if len(non_empty) >= 3:
                    match_count = sum(1 for v in norm_vals if any(k in v for k in ["họ", "tên", "email", "masv", "mssv", "mã", "sđt", "điện thoại", "lớp", "chuyên ngành"]))
                    if match_count >= 2:
                        header_row_idx = r
                        header_vals = row_vals
                        break

            if header_row_idx is None:
                header_row_idx = 4 if ws.max_row >= 4 else 1
                header_vals = [str(ws.cell(row=header_row_idx, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]

            col_map = _detect_column_mapping(header_vals)

            for r in range(header_row_idx + 1, ws.max_row + 1):
                row_cells = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                str_cells = [str(v).strip() if v is not None else "" for v in row_cells]
                if not any(str_cells):
                    continue

                def _get_val(key):
                    if key in col_map and col_map[key] < len(str_cells):
                        return str_cells[col_map[key]]
                    return ""

                raw_rows.append({
                    "dong": r,
                    "ma_sinh_vien": _get_val("ma_sinh_vien"),
                    "ho_ten": _get_val("ho_ten"),
                    "email": _get_val("email"),
                    "so_dien_thoai": _get_val("so_dien_thoai"),
                    "lop": _get_val("lop"),
                    "chuyen_nganh": _get_val("chuyen_nganh"),
                    "truong": _get_val("truong"),
                })
        except Exception as e:
            file_errors.append(f"Không thể đọc file Excel: {str(e)}")

    elif ext == "docx":
        if docx is None:
            file_errors.append("Hệ thống chưa cài đặt thư viện python-docx để xử lý file Word.")
        else:
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                if not doc.tables:
                    file_errors.append("File Word không chứa bảng dữ liệu (Table) danh sách sinh viên.")
                else:
                    table = doc.tables[0]
                    rows = table.rows
                    if len(rows) < 2:
                        file_errors.append("Bảng trong file Word không có dữ liệu sinh viên.")
                    else:
                        header_vals = [cell.text.strip() for cell in rows[0].cells]
                        col_map = _detect_column_mapping(header_vals)

                        for r_idx, row in enumerate(rows[1:], 2):
                            str_cells = [cell.text.strip() for cell in row.cells]
                            if not any(str_cells):
                                continue

                            def _get_val(key):
                                if key in col_map and col_map[key] < len(str_cells):
                                    return str_cells[col_map[key]]
                                return ""

                            raw_rows.append({
                                "dong": r_idx,
                                "ma_sinh_vien": _get_val("ma_sinh_vien"),
                                "ho_ten": _get_val("ho_ten"),
                                "email": _get_val("email"),
                                "so_dien_thoai": _get_val("so_dien_thoai"),
                                "lop": _get_val("lop"),
                                "chuyen_nganh": _get_val("chuyen_nganh"),
                                "truong": _get_val("truong"),
                            })
            except Exception as e:
                file_errors.append(f"Không thể đọc file Word .docx: {str(e)}")

    elif ext == "csv":
        try:
            content = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(content))
            all_rows = list(reader)
            if not all_rows:
                file_errors.append("File CSV trống.")
            else:
                header_vals = all_rows[0]
                col_map = _detect_column_mapping(header_vals)
                for r_idx, row in enumerate(all_rows[1:], 2):
                    str_cells = [v.strip() for v in row]
                    if not any(str_cells):
                        continue

                    def _get_val(key):
                        if key in col_map and col_map[key] < len(str_cells):
                            return str_cells[col_map[key]]
                        return ""

                    raw_rows.append({
                        "dong": r_idx,
                        "ma_sinh_vien": _get_val("ma_sinh_vien"),
                        "ho_ten": _get_val("ho_ten"),
                        "email": _get_val("email"),
                        "so_dien_thoai": _get_val("so_dien_thoai"),
                        "lop": _get_val("lop"),
                        "chuyen_nganh": _get_val("chuyen_nganh"),
                        "truong": _get_val("truong"),
                    })
        except Exception as e:
            file_errors.append(f"Không thể đọc file CSV: {str(e)}")
    elif ext == "pdf":
        if pypdf is None:
            file_errors.append("Hệ thống chưa cài đặt thư viện pypdf để xử lý file PDF.")
        else:
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                if len(reader.pages) == 0:
                    file_errors.append("File PDF không chứa trang dữ liệu nào.")
                else:
                    lines = []
                    for page in reader.pages:
                        txt = page.extract_text() or ""
                        for line in txt.splitlines():
                            l = line.strip()
                            if l:
                                lines.append(l)

                    if not lines:
                        file_errors.append("File PDF không chứa nội dung văn bản có thể đọc được.")
                    else:
                        header_row_idx = None
                        delimiter = None
                        for idx, line in enumerate(lines[:15]):
                            norm = _normalize_header(line)
                            if any(k in norm for k in ["họ", "tên", "email", "masv", "mssv", "mã"]):
                                header_row_idx = idx
                                if "|" in line:
                                    delimiter = "|"
                                elif "\t" in line:
                                    delimiter = "\t"
                                else:
                                    delimiter = "space"
                                break

                        if header_row_idx is not None:
                            if delimiter == "|":
                                header_vals = [c.strip() for c in lines[header_row_idx].split("|") if c.strip()]
                            elif delimiter == "\t":
                                header_vals = [c.strip() for c in lines[header_row_idx].split("\t") if c.strip()]
                            else:
                                header_vals = [c.strip() for c in re.split(r"\s{2,}", lines[header_row_idx]) if c.strip()]

                            col_map = _detect_column_mapping(header_vals)

                            for r_idx, line in enumerate(lines[header_row_idx + 1:], header_row_idx + 2):
                                if delimiter == "|":
                                    str_cells = [c.strip() for c in line.split("|") if c.strip()]
                                elif delimiter == "\t":
                                    str_cells = [c.strip() for c in line.split("\t") if c.strip()]
                                else:
                                    str_cells = [c.strip() for c in re.split(r"\s{2,}", line) if c.strip()]

                                if not any(str_cells) or len(str_cells) < 2:
                                    continue

                                def _get_val(key):
                                    if key in col_map and col_map[key] < len(str_cells):
                                        return str_cells[col_map[key]]
                                    return ""

                                raw_rows.append({
                                    "dong": r_idx,
                                    "ma_sinh_vien": _get_val("ma_sinh_vien"),
                                    "ho_ten": _get_val("ho_ten"),
                                    "email": _get_val("email"),
                                    "so_dien_thoai": _get_val("so_dien_thoai"),
                                    "lop": _get_val("lop"),
                                    "chuyen_nganh": _get_val("chuyen_nganh"),
                                    "truong": _get_val("truong"),
                                })

                        if not raw_rows:
                            email_indices = [i for i, l in enumerate(lines) if "@" in l and "." in l and " " not in l]
                            if email_indices:
                                for idx_ord, e_idx in enumerate(email_indices, 1):
                                    email = lines[e_idx]
                                    ho_ten = lines[e_idx - 1] if e_idx - 1 >= 0 else ""
                                    ma_sv = lines[e_idx - 2] if e_idx - 2 >= 0 else ""
                                    if ma_sv.isdigit() and e_idx - 3 >= 0:
                                        ma_sv = lines[e_idx - 1]
                                        ho_ten = lines[e_idx - 2]
                                    sdt = lines[e_idx + 1] if e_idx + 1 < len(lines) else ""
                                    lop = lines[e_idx + 2] if e_idx + 2 < len(lines) else ""
                                    chuyen_nganh = lines[e_idx + 3] if e_idx + 3 < len(lines) else ""
                                    next_limit = email_indices[idx_ord] - 2 if idx_ord < len(email_indices) else len(lines)
                                    truong_parts = []
                                    for t_idx in range(e_idx + 4, min(next_limit, len(lines))):
                                        if lines[t_idx].isdigit() and int(lines[t_idx]) == idx_ord + 1:
                                            break
                                        truong_parts.append(lines[t_idx])
                                    truong = " ".join(truong_parts)
                                    raw_rows.append({
                                        "dong": idx_ord + 1,
                                        "ma_sinh_vien": ma_sv,
                                        "ho_ten": ho_ten,
                                        "email": email,
                                        "so_dien_thoai": sdt,
                                        "lop": lop,
                                        "chuyen_nganh": chuyen_nganh,
                                        "truong": truong,
                                    })
                        if not raw_rows:
                            for r_idx, line in enumerate(lines, 1):
                                parts = [p.strip() for p in re.split(r"[\t|]+|\s{2,}", line) if p.strip()]
                                if len(parts) >= 2:
                                    email_idx = next((i for i, p in enumerate(parts) if "@" in p), -1)
                                    if email_idx != -1:
                                        raw_rows.append({
                                            "dong": r_idx,
                                            "ma_sinh_vien": parts[0] if email_idx > 0 else "",
                                            "ho_ten": parts[1] if len(parts) > 1 and email_idx != 1 else parts[0],
                                            "email": parts[email_idx],
                                            "so_dien_thoai": parts[email_idx + 1] if len(parts) > email_idx + 1 else "",
                                            "lop": "",
                                            "chuyen_nganh": "",
                                            "truong": "",
                                        })
            except Exception as e:
                file_errors.append(f"Không thể đọc file PDF: {str(e)}")
    else:
        file_errors.append(f"Định dạng file .{ext} không được hỗ trợ. Vui lòng sử dụng file .xlsx, .xls, .docx, .pdf hoặc .csv.")

    return raw_rows, file_errors
