import io
import os
from datetime import datetime
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


def _register_pdf_fonts():
    """
    Đăng ký font Unicode tiếng Việt cho ReportLab.
    Ưu tiên lấy font Arial trên hệ điều hành Windows để hỗ trợ đầy đủ tiếng Việt có dấu.
    Nếu không tìm thấy hoặc lỗi sẽ fallback về Helvetica.
    """
    regular_font = "Helvetica"
    bold_font = "Helvetica-Bold"

    windows_fonts = {
        "regular": "C:/Windows/Fonts/arial.ttf",
        "bold": "C:/Windows/Fonts/arialbd.ttf"
    }

    try:
        if os.path.exists(windows_fonts["regular"]):
            pdfmetrics.registerFont(TTFont("VN-Arial", windows_fonts["regular"]))
            regular_font = "VN-Arial"

            if os.path.exists(windows_fonts["bold"]):
                pdfmetrics.registerFont(TTFont("VN-Arial-Bold", windows_fonts["bold"]))
                bold_font = "VN-Arial-Bold"
            else:
                bold_font = "VN-Arial"
    except Exception:
        regular_font = "Helvetica"
        bold_font = "Helvetica-Bold"

    return regular_font, bold_font


def export_evaluations_to_excel(summary_data: Dict[str, Any], items: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Xuất báo cáo tổng hợp và danh sách đánh giá ra file Excel (.xlsx).
    Bao gồm 2 sheet:
    - Sheet 1: Thống kê chỉ số KPI, phân bổ xếp loại và điểm trung bình theo trường.
    - Sheet 2: Danh sách chi tiết từng bản đánh giá của thực tập sinh.
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Tổng hợp KPI
    ws_kpi = wb.active
    ws_kpi.title = "Tong_Hop_KPI"
    ws_kpi.views.sheetView[0].showGridLines = True

    # Định nghĩa styles
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1F497D")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    section_font = Font(name="Segoe UI", size=12, bold=True, color="1F497D")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=11)
    bold_data_font = Font(name="Segoe UI", size=11, bold=True)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # 1. TIÊU ĐỀ BÁO CÁO
    ws_kpi.merge_cells("A1:E1")
    ws_kpi["A1"] = "BÁO CÁO TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ THỰC TẬP SINH"
    ws_kpi["A1"].font = title_font
    ws_kpi["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_kpi.row_dimensions[1].height = 30

    ws_kpi.merge_cells("A2:E2")
    ws_kpi["A2"] = f"Ngày xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws_kpi["A2"].font = subtitle_font
    ws_kpi["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # 2. KHỐI CHỈ SỐ KPI CHUNG
    ws_kpi["A4"] = "I. CHỈ SỐ TỔNG HỢP CHUNG (KPIs)"
    ws_kpi["A4"].font = section_font

    kpi_headers = ["Chỉ số đánh giá", "Giá trị"]
    for col_idx, h in enumerate(kpi_headers, 1):
        cell = ws_kpi.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_kpi.row_dimensions[5].height = 24

    kpi_rows = [
        ("Tổng số hồ sơ thực tập sinh", summary_data.get("tong_so_ho_so", 0)),
        ("Tổng số lượt đánh giá", summary_data.get("tong_so_danh_gia", 0)),
        ("Điểm kỹ năng chuyên môn trung bình", f"{summary_data.get('diem_ky_nang_tb', 0):.2f} / 10"),
        ("Điểm thái độ kỷ luật trung bình", f"{summary_data.get('diem_thai_do_tb', 0):.2f} / 10"),
        ("Điểm tổng kết trung bình toàn đoàn", f"{summary_data.get('diem_tong_ket_tb', 0):.2f} / 10"),
        ("Số lượng đề xuất tuyển dụng chính thức", f"{summary_data.get('so_luong_de_xuat_tuyen_dung', 0)} thực tập sinh"),
        ("Tỷ lệ đề xuất tuyển dụng chính thức", f"{summary_data.get('ty_le_de_xuat_tuyen_dung', 0):.1f}%"),
    ]

    for idx, (label, val) in enumerate(kpi_rows, 6):
        c1 = ws_kpi.cell(row=idx, column=1, value=label)
        c2 = ws_kpi.cell(row=idx, column=2, value=val)
        c1.font = data_font
        c2.font = bold_data_font
        c1.border = thin_border
        c2.border = thin_border
        c2.alignment = Alignment(horizontal="center")
        if idx % 2 == 0:
            c1.fill = alt_row_fill
            c2.fill = alt_row_fill

    # 3. PHÂN BỔ XẾP LOẠI
    start_r = 15
    ws_kpi.cell(row=start_r, column=1, value="II. PHÂN BỔ THEO XẾP LOẠI RÈN LUYỆN").font = section_font
    dist_headers = ["Xếp loại", "Tiêu chuẩn điểm", "Số lượng", "Tỷ lệ (%)"]
    for col_idx, h in enumerate(dist_headers, 1):
        cell = ws_kpi.cell(row=start_r + 1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_kpi.row_dimensions[start_r + 1].height = 24

    phan_bo = summary_data.get("phan_bo_xep_loai", {})
    total_dg = summary_data.get("tong_so_danh_gia", 0) or 1
    grade_rows = [
        ("Xuất sắc", "Điểm TB >= 9.0", phan_bo.get("XuatSac", 0)),
        ("Giỏi", "8.0 <= Điểm TB < 9.0", phan_bo.get("Gioi", 0)),
        ("Khá", "6.5 <= Điểm TB < 8.0", phan_bo.get("Kha", 0)),
        ("Trung bình", "5.0 <= Điểm TB < 6.5", phan_bo.get("TrungBinh", 0)),
        ("Yếu / Không đạt", "Điểm TB < 5.0", phan_bo.get("Yeu", 0)),
    ]

    for idx, (gr, crit, cnt) in enumerate(grade_rows, start_r + 2):
        ratio = (cnt / total_dg) * 100.0 if total_dg else 0
        c1 = ws_kpi.cell(row=idx, column=1, value=gr)
        c2 = ws_kpi.cell(row=idx, column=2, value=crit)
        c3 = ws_kpi.cell(row=idx, column=3, value=cnt)
        c4 = ws_kpi.cell(row=idx, column=4, value=f"{ratio:.1f}%")
        for c in (c1, c2, c3, c4):
            c.font = data_font
            c.border = thin_border
        c3.alignment = Alignment(horizontal="center")
        c4.alignment = Alignment(horizontal="center")
        if idx % 2 == 0:
            for c in (c1, c2, c3, c4):
                c.fill = alt_row_fill

    # 4. THỐNG KÊ THEO TRƯỜNG ĐẠI HỌC
    start_r2 = start_r + len(grade_rows) + 3
    ws_kpi.cell(row=start_r2, column=1, value="III. THỐNG KÊ THEO TRƯỜNG ĐẠI HỌC").font = section_font
    uni_headers = ["Mã trường", "Tên trường đại học", "Số lượng đánh giá", "Điểm trung bình"]
    for col_idx, h in enumerate(uni_headers, 1):
        cell = ws_kpi.cell(row=start_r2 + 1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_kpi.row_dimensions[start_r2 + 1].height = 24

    uni_list = summary_data.get("thong_ke_theo_truong", [])
    if not uni_list:
        c = ws_kpi.cell(row=start_r2 + 2, column=1, value="Chưa có dữ liệu theo trường")
        ws_kpi.merge_cells(start_row=start_r2 + 2, start_column=1, end_row=start_r2 + 2, end_column=4)
        c.alignment = Alignment(horizontal="center")
        c.font = data_font
    else:
        for idx, u in enumerate(uni_list, start_r2 + 2):
            c1 = ws_kpi.cell(row=idx, column=1, value=u.get("ma_truong") or "-")
            c2 = ws_kpi.cell(row=idx, column=2, value=u.get("ten_truong", ""))
            c3 = ws_kpi.cell(row=idx, column=3, value=u.get("so_luong_danh_gia", 0))
            c4 = ws_kpi.cell(row=idx, column=4, value=f"{u.get('diem_trung_binh', 0):.2f}")
            for c in (c1, c2, c3, c4):
                c.font = data_font
                c.border = thin_border
            c1.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            c4.alignment = Alignment(horizontal="center")
            if idx % 2 == 0:
                for c in (c1, c2, c3, c4):
                    c.fill = alt_row_fill

    # Tự động co giãn độ rộng cột Sheet 1
    for col in ws_kpi.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and not cell.coordinate in ["A1", "A2"]:
                max_len = len(val_str)
        ws_kpi.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # -------------------------------------------------------------
    # Sheet 2: Danh sách chi tiết đánh giá
    # -------------------------------------------------------------
    ws_detail = wb.create_sheet(title="Chi_Tiet_Danh_Gia")
    ws_detail.views.sheetView[0].showGridLines = True

    detail_headers = [
        "STT",
        "Mã ĐG",
        "Mã HS",
        "Họ và tên",
        "Email",
        "Số điện thoại",
        "Chuyên ngành",
        "Trường đại học",
        "Trạng thái thực tập",
        "Đợt đánh giá",
        "Điểm kỹ năng",
        "Điểm thái độ",
        "Điểm TB",
        "Xếp loại",
        "Đề xuất tuyển dụng",
        "Người đánh giá",
        "Nhận xét chi tiết"
    ]

    ws_detail.row_dimensions[1].height = 28
    for col_idx, h in enumerate(detail_headers, 1):
        cell = ws_detail.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for idx, item in enumerate(items, 2):
        ws_detail.row_dimensions[idx].height = 20
        row_vals = [
            idx - 1,
            item.get("ma_danh_gia"),
            item.get("ma_ho_so"),
            item.get("ho_ten") or "",
            item.get("email") or "",
            item.get("so_dien_thoai") or "",
            item.get("chuyen_nganh") or "",
            item.get("ten_truong") or "",
            item.get("trang_thai_thuc_tap") or "",
            item.get("loai_danh_gia") or "",
            f"{item.get('diem_ky_nang', 0):.1f}",
            f"{item.get('diem_thai_do', 0):.1f}",
            f"{item.get('diem_trung_binh', 0):.2f}",
            item.get("xep_loai") or "",
            "Có" if item.get("de_xuat_tuyen_chinh_thuc") else "Không",
            item.get("nguoi_danh_gia") or "",
            item.get("nhan_xet_chi_tiet") or ""
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_detail.cell(row=idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            # Căn lề số/mã ở giữa
            if col_idx in [1, 2, 3, 6, 9, 10, 11, 12, 13, 14, 15]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if idx % 2 == 1:
                cell.fill = alt_row_fill

    # Tự động co giãn độ rộng cột Sheet 2
    for col in ws_detail.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_detail.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_evaluations_to_pdf(summary_data: Dict[str, Any], items: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Xuất báo cáo tổng hợp và danh sách đánh giá ra file PDF.
    Cấu hình khổ giấy A4 ngang (Landscape) để hiển thị đầy đủ thông tin,
    sử dụng font Unicode tiếng Việt chuẩn không bị lỗi ký tự có dấu.
    """
    regular_font, bold_font = _register_pdf_fonts()
    output = io.BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=25,
        rightMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        name="ReportTitle",
        fontName=bold_font,
        fontSize=16,
        leading=20,
        alignment=1, # Center
        textColor=colors.HexColor("#1F4E78")
    )

    subtitle_style = ParagraphStyle(
        name="ReportSubtitle",
        fontName=regular_font,
        fontSize=9,
        leading=12,
        alignment=1, # Center
        textColor=colors.HexColor("#595959")
    )

    heading2_style = ParagraphStyle(
        name="ReportHeading2",
        fontName=bold_font,
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1F4E78"),
        spaceBefore=10,
        spaceAfter=5
    )

    body_style = ParagraphStyle(
        name="ReportBody",
        fontName=regular_font,
        fontSize=8.5,
        leading=11,
    )

    body_center_style = ParagraphStyle(
        name="ReportBodyCenter",
        fontName=regular_font,
        fontSize=8.5,
        leading=11,
        alignment=1,
    )

    header_cell_style = ParagraphStyle(
        name="ReportHeaderCell",
        fontName=bold_font,
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.white
    )

    elements = []

    # 1. TIÊU ĐỀ
    elements.append(Paragraph("BÁO CÁO TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ THỰC TẬP SINH", title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(f"Thời điểm xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", subtitle_style))
    elements.append(Spacer(1, 12))

    # 2. BẢNG KPI TỔNG HỢP VÀ PHÂN BỔ XẾP LOẠI (Hiển thị 2 cột cạnh nhau)
    elements.append(Paragraph("I. CHỈ SỐ KPI TỔNG QUAN & PHÂN BỔ XẾP LOẠI", heading2_style))

    phan_bo = summary_data.get("phan_bo_xep_loai", {})
    total_dg = summary_data.get("tong_so_danh_gia", 0) or 1

    # Bảng 1: KPI
    kpi_table_data = [
        [Paragraph("Chỉ số đánh giá", header_cell_style), Paragraph("Giá trị", header_cell_style)],
        [Paragraph("Tổng số hồ sơ thực tập sinh", body_style), Paragraph(str(summary_data.get("tong_so_ho_so", 0)), body_center_style)],
        [Paragraph("Tổng số lượt đánh giá", body_style), Paragraph(str(summary_data.get("tong_so_danh_gia", 0)), body_center_style)],
        [Paragraph("Điểm kỹ năng chuyên môn TB", body_style), Paragraph(f"{summary_data.get('diem_ky_nang_tb', 0):.2f} / 10", body_center_style)],
        [Paragraph("Điểm thái độ kỷ luật TB", body_style), Paragraph(f"{summary_data.get('diem_thai_do_tb', 0):.2f} / 10", body_center_style)],
        [Paragraph("Điểm tổng kết TB toàn đoàn", body_style), Paragraph(f"{summary_data.get('diem_tong_ket_tb', 0):.2f} / 10", body_center_style)],
        [Paragraph("Đề xuất tuyển dụng chính thức", body_style), Paragraph(f"{summary_data.get('so_luong_de_xuat_tuyen_dung', 0)} ({summary_data.get('ty_le_de_xuat_tuyen_dung', 0):.1f}%)", body_center_style)],
    ]

    t_kpi = Table(kpi_table_data, colWidths=[240, 140])
    t_kpi.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    # Bảng 2: Phân bổ xếp loại
    grade_table_data = [
        [Paragraph("Xếp loại", header_cell_style), Paragraph("Tiêu chuẩn", header_cell_style), Paragraph("Số lượng", header_cell_style), Paragraph("Tỷ lệ", header_cell_style)],
        [Paragraph("Xuất sắc", body_style), Paragraph(">= 9.0", body_center_style), Paragraph(str(phan_bo.get("XuatSac", 0)), body_center_style), Paragraph(f"{(phan_bo.get('XuatSac', 0)/total_dg)*100:.1f}%", body_center_style)],
        [Paragraph("Giỏi", body_style), Paragraph("8.0 - < 9.0", body_center_style), Paragraph(str(phan_bo.get("Gioi", 0)), body_center_style), Paragraph(f"{(phan_bo.get('Gioi', 0)/total_dg)*100:.1f}%", body_center_style)],
        [Paragraph("Khá", body_style), Paragraph("6.5 - < 8.0", body_center_style), Paragraph(str(phan_bo.get("Kha", 0)), body_center_style), Paragraph(f"{(phan_bo.get('Kha', 0)/total_dg)*100:.1f}%", body_center_style)],
        [Paragraph("Trung bình", body_style), Paragraph("5.0 - < 6.5", body_center_style), Paragraph(str(phan_bo.get("TrungBinh", 0)), body_center_style), Paragraph(f"{(phan_bo.get('TrungBinh', 0)/total_dg)*100:.1f}%", body_center_style)],
        [Paragraph("Yếu / KĐ", body_style), Paragraph("< 5.0", body_center_style), Paragraph(str(phan_bo.get("Yeu", 0)), body_center_style), Paragraph(f"{(phan_bo.get('Yeu', 0)/total_dg)*100:.1f}%", body_center_style)],
    ]

    t_grade = Table(grade_table_data, colWidths=[100, 100, 70, 80])
    t_grade.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    # Ghép 2 bảng nằm ngang trong 1 Table cha
    layout_table = Table([[t_kpi, t_grade]], colWidths=[400, 370])
    layout_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    elements.append(layout_table)
    elements.append(Spacer(1, 10))

    # 3. BẢNG THỐNG KÊ THEO TRƯỜNG ĐẠI HỌC
    elements.append(Paragraph("II. THỐNG KÊ KẾT QUẢ THEO TRƯỜNG ĐẠI HỌC", heading2_style))
    uni_list = summary_data.get("thong_ke_theo_truong", [])
    uni_table_data = [
        [
            Paragraph("Mã trường", header_cell_style),
            Paragraph("Tên trường đại học", header_cell_style),
            Paragraph("Số lượng đánh giá", header_cell_style),
            Paragraph("Điểm trung bình", header_cell_style),
        ]
    ]

    if uni_list:
        for u in uni_list:
            uni_table_data.append([
                Paragraph(str(u.get("ma_truong") or "-"), body_center_style),
                Paragraph(u.get("ten_truong", ""), body_style),
                Paragraph(str(u.get("so_luong_danh_gia", 0)), body_center_style),
                Paragraph(f"{u.get('diem_trung_binh', 0):.2f}", body_center_style),
            ])
    else:
        uni_table_data.append([
            Paragraph("-", body_center_style),
            Paragraph("Không có dữ liệu", body_style),
            Paragraph("0", body_center_style),
            Paragraph("0.00", body_center_style),
        ])

    t_uni = Table(uni_table_data, colWidths=[80, 420, 140, 130])
    t_uni.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(t_uni)
    elements.append(Spacer(1, 10))

    # 4. DANH SÁCH CHI TIẾT ĐÁNH GIÁ CỦA CÁC THỰC TẬP SINH
    elements.append(Paragraph("III. DANH SÁCH CHI TIẾT ĐÁNH GIÁ", heading2_style))
    detail_table_data = [
        [
            Paragraph("STT", header_cell_style),
            Paragraph("Họ và tên", header_cell_style),
            Paragraph("Trường ĐH", header_cell_style),
            Paragraph("Chuyên ngành", header_cell_style),
            Paragraph("Đợt", header_cell_style),
            Paragraph("Kỹ năng", header_cell_style),
            Paragraph("Thái độ", header_cell_style),
            Paragraph("Điểm TB", header_cell_style),
            Paragraph("Xếp loại", header_cell_style),
            Paragraph("Đề xuất TD", header_cell_style),
            Paragraph("Nhận xét", header_cell_style),
        ]
    ]

    for idx, it in enumerate(items, 1):
        detail_table_data.append([
            Paragraph(str(idx), body_center_style),
            Paragraph(it.get("ho_ten") or "", body_style),
            Paragraph(it.get("ten_truong") or "", body_style),
            Paragraph(it.get("chuyen_nganh") or "", body_style),
            Paragraph(it.get("loai_danh_gia") or "", body_center_style),
            Paragraph(f"{it.get('diem_ky_nang', 0):.1f}", body_center_style),
            Paragraph(f"{it.get('diem_thai_do', 0):.1f}", body_center_style),
            Paragraph(f"{it.get('diem_trung_binh', 0):.2f}", body_center_style),
            Paragraph(it.get("xep_loai") or "", body_center_style),
            Paragraph("Có" if it.get("de_xuat_tuyen_chinh_thuc") else "Không", body_center_style),
            Paragraph(it.get("nhan_xet_chi_tiet") or "", body_style),
        ])

    # Tổng chiều rộng: 30 + 100 + 110 + 90 + 45 + 45 + 45 + 50 + 55 + 60 + 140 = 770 (vừa vặn landscape A4)
    t_detail = Table(
        detail_table_data,
        colWidths=[30, 100, 110, 90, 45, 45, 45, 50, 55, 60, 140]
    )
    t_detail.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(t_detail)

    doc.build(elements)
    output.seek(0)
    return output
