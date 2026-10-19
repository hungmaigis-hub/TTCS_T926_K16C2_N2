"""
Script sinh báo cáo kiểm thử Excel Sprint 4 chuyên nghiệp cho Tester: Trần Quốc Huy
File xuất: Bao_Cao_Kiem_Thu_Sprint4_Tran_Quoc_Huy.xlsx
Bao gồm 3 Sheet:
  1. Tổng Quan Sprint 4 (KPI, Bảng tóm tắt 3 nhiệm vụ)
  2. Chi Tiết 54 Test Cases (Đầy đủ 54 TCs, 100% PASSED)
  3. Báo Cáo Lỗi & Đề Xuất (Các ghi nhận, phân tích ngoại lệ và khuyến nghị nghiệp vụ)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
# Xóa sheet mặc định
wb.remove(wb.active)

# ==========================================
# CẤU HÌNH PHÔNG CHỮ VÀ BẢNG MÀU CHUẨN
# ==========================================
font_title = Font(name="Calibri", size=16, bold=True, color="1F4E78")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
font_section = Font(name="Calibri", size=13, bold=True, color="1F4E78")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=10, color="000000")
font_data_bold = Font(name="Calibri", size=10, bold=True, color="000000")
font_pass = Font(name="Calibri", size=10, bold=True, color="276A3C")
font_kpi_num = Font(name="Calibri", size=18, bold=True, color="1F4E78")
font_kpi_pass = Font(name="Calibri", size=18, bold=True, color="276A3C")
font_kpi_fail = Font(name="Calibri", size=18, bold=True, color="C00000")

fill_navy_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fill_blue_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
fill_pass = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fill_kpi_box = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

thin_gray = Side(border_style="thin", color="D9D9D9")
med_navy = Side(border_style="medium", color="1F4E78")
double_navy = Side(border_style="double", color="1F4E78")

border_cell = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
border_header = Border(left=thin_gray, right=thin_gray, top=med_navy, bottom=med_navy)
border_total = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=double_navy)

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ==============================================================================
# SHEET 1: TỔNG QUAN SPRINT 4
# ==============================================================================
ws1 = wb.create_sheet(title="Tổng Quan Sprint 4")
ws1.views.sheetView[0].showGridLines = True

# Header Title
ws1.cell(2, 2, "BÁO CÁO KẾT QUẢ KIỂM THỬ SPRINT 4 (TEST EXECUTION REPORT)").font = font_title
ws1.cell(3, 2, "Hệ thống Quản lý Thực tập sinh | Thành viên phụ trách: Trần Quốc Huy (Tester)").font = font_subtitle

# Bảng thông tin dự án (B5:C9)
proj_info = [
    ("Dự án / Học phần:", "Thực tập cơ sở (TTCS) - Quản lý Thực tập sinh"),
    ("Giai đoạn (Sprint):", "Sprint 4"),
    ("Tester phụ trách:", "Trần Quốc Huy"),
    ("Công cụ kiểm thử:", "Pytest, FastAPI TestClient, SQLite Isolated Engine"),
    ("Ngày thực hiện & xuất báo cáo:", "19/10/2026")
]
for idx, (label, val) in enumerate(proj_info, start=5):
    c_lbl = ws1.cell(idx, 2, label)
    c_lbl.font = font_data_bold
    c_lbl.fill = fill_blue_sub
    c_lbl.border = border_cell
    c_val = ws1.cell(idx, 3, val)
    c_val.font = font_data
    c_val.border = border_cell

# Bảng KPI Chỉ số kiểm thử (E5:G10)
kpi_headers = ["CHỈ SỐ KIỂM THỬ", "SỐ LƯỢNG", "TỶ LỆ (%)"]
for c_idx, h in enumerate(kpi_headers, start=5):
    cell = ws1.cell(5, c_idx, h)
    cell.font = font_header
    cell.fill = fill_navy_header
    cell.alignment = align_center
    cell.border = border_header

kpi_rows = [
    ("Tổng số User Story phụ trách", "3", "-"),
    ("Tổng số Test Case thực thi", "54", "100.0%"),
    ("Số Test Case ĐẠT (PASSED)", "54", "100.0%"),
    ("Số Test Case KHÔNG ĐẠT (FAILED)", "0", "0.0%"),
    ("Tổng số Bug phát hiện & bàn giao Dev", "0", "0.0%")
]
for r_idx, (kpi_name, qty, rate) in enumerate(kpi_rows, start=6):
    c1 = ws1.cell(r_idx, 5, kpi_name)
    c1.font = font_data
    c1.border = border_cell
    
    c2 = ws1.cell(r_idx, 6, int(qty) if qty.isdigit() else qty)
    c2.font = font_data_bold if r_idx in (7, 8) else font_data
    if r_idx == 8: c2.font = font_pass
    c2.alignment = align_center
    c2.border = border_cell

    c3 = ws1.cell(r_idx, 7, rate)
    c3.font = font_pass if r_idx in (7, 8) else font_data
    c3.alignment = align_center
    c3.border = border_cell

# Bảng tóm tắt 3 User Story kiểm thử (Row 12 -> 16)
ws1.cell(11, 2, "BẢNG TỔNG HỢP KIỂM THỬ THEO TỪNG NHIỆM VỤ BACKLOG SPRINT 4").font = font_section

story_headers = ["STT", "Test Key", "Dòng Backlog", "Chức năng / Nhiệm vụ kiểm thử", "Số TC", "Passed", "Failed", "Đánh giá sơ bộ"]
for c_idx, h in enumerate(story_headers, start=2):
    cell = ws1.cell(12, c_idx, h)
    cell.font = font_header
    cell.fill = fill_navy_header
    cell.alignment = align_center
    cell.border = border_header

stories_data = [
    (1, "TC-AttendanceScan", "Nhiệm vụ 1", "Là thực tập sinh, quét mã QR / thẻ RFID để chấm công vào ca nhanh chóng", 18, 18, 0, "Hoàn thành 100%, phân loại Đúng Giờ / Đi Muộn / Về Sớm, phân giải prefix INTERN/HS/email/SĐT, chống quét đúp < 60s và chặn quét lần 3"),
    (2, "TC-Meetings", "Nhiệm vụ 2", "Là Mentor / HR, tạo lịch họp và gửi email HTML thông báo tự động cho sinh viên", 18, 18, 0, "Hoàn thành 100%, tích hợp BackgroundTasks gửi email HTML, lọc trùng ID, validate thời gian bắt đầu/kết thúc và kiểm soát tài khoản khóa"),
    (3, "TC-SystemBackup", "Nhiệm vụ 3", "Là Admin, sao lưu dữ liệu tự động định kỳ để đảm bảo an toàn thông tin hệ thống", 18, 18, 0, "Hoàn thành 100%, nén gzip .sql.gz an toàn, mã băm SHA-256 toàn vẹn, chống tấn công Path Traversal (../) và Concurrency Lock (409)")
]

for r_idx, row in enumerate(stories_data, start=13):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate(row, start=2):
        cell = ws1.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = border_cell
        if fill.fill_type: cell.fill = fill
        if c_idx in (2, 3, 4, 6, 7, 8): cell.alignment = align_center
        else: cell.alignment = align_left
        if c_idx == 7 and val > 0:
            cell.font = font_pass

# Dòng tổng cộng
tot_row = 16
ws1.cell(tot_row, 2, "").border = border_total
ws1.cell(tot_row, 3, "TỔNG CỘNG").font = font_data_bold
ws1.cell(tot_row, 3).alignment = align_center
ws1.cell(tot_row, 3).border = border_total
ws1.cell(tot_row, 4, "Sprint 4").alignment = align_center
ws1.cell(tot_row, 4).border = border_total
ws1.cell(tot_row, 5, "Hoàn thành toàn diện 3 User Story theo kế hoạch Backlog").font = font_data_bold
ws1.cell(tot_row, 5).border = border_total
c_tot = ws1.cell(tot_row, 6, 54)
c_tot.font = font_data_bold
c_tot.alignment = align_center
c_tot.border = border_total
c_pass = ws1.cell(tot_row, 7, 54)
c_pass.font = font_pass
c_pass.alignment = align_center
c_pass.border = border_total
c_fail = ws1.cell(tot_row, 8, 0)
c_fail.font = font_data_bold
c_fail.alignment = align_center
c_fail.border = border_total
c_eval = ws1.cell(tot_row, 9, "100% PASSED (54/54 Test Cases đạt chuẩn chất lượng)")
c_eval.font = font_pass
c_eval.border = border_total

# Độ rộng cột Sheet 1
ws1.column_dimensions['A'].width = 4
ws1.column_dimensions['B'].width = 8
ws1.column_dimensions['C'].width = 24
ws1.column_dimensions['D'].width = 16
ws1.column_dimensions['E'].width = 45
ws1.column_dimensions['F'].width = 12
ws1.column_dimensions['G'].width = 14
ws1.column_dimensions['H'].width = 12
ws1.column_dimensions['I'].width = 45


# ==============================================================================
# SHEET 2: CHI TIẾT 54 TEST CASES
# ==============================================================================
ws2 = wb.create_sheet(title="Chi Tiết 54 Test Cases")
ws2.views.sheetView[0].showGridLines = True

tc_headers = [
    "STT", "Test Key", "Test Case ID", "User Story",
    "Mục tiêu / Tiêu đề kịch bản kiểm thử", "Loại kiểm thử",
    "Dữ liệu đầu vào (Test Data / Payload)",
    "Kết quả kỳ vọng (Expected Result)",
    "Kết quả thực tế (Actual Result)",
    "Trạng thái (Status)", "Ghi chú (Notes)"
]

for c_idx, h in enumerate(tc_headers, start=1):
    cell = ws2.cell(1, c_idx, h)
    cell.font = font_header
    cell.fill = fill_navy_header
    cell.alignment = align_center
    cell.border = border_header

tc_54_data = [
    # --- NHIỆM VỤ 1: CHẤM CÔNG QUÉT MÃ QR / CARD RFID (18 TCs) ---
    (1, "TC-AttendanceScan", "TC_ATT_01", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Quét mã QR check-in đúng giờ (< 08:30)", "Functional",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "1", "phuong_thuc": "QR", "thoi_gian_quet": "08:15:00"}',
     "Status 200 OK, hanh_dong = 'CHECK_IN', trang_thai = 'DungGio'",
     "Status 200 OK, hanh_dong = CHECK_IN, trang_thai = DungGio", "PASSED", "DungGio trước 08:30"),

    (2, "TC-AttendanceScan", "TC_ATT_02", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Quét mã thẻ Card RFID check-in đúng giờ", "Functional",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "CARD-001", "phuong_thuc": "Card", "thoi_gian_quet": "08:20:00"}',
     "Status 200 OK, phuong_thuc = 'Card', trang_thai = 'DungGio'",
     "Status 200 OK, phuong_thuc = Card, trang_thai = DungGio", "PASSED", "Thẻ RFID hợp lệ"),

    (3, "TC-AttendanceScan", "TC_ATT_03", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Quét check-in đi muộn (> 08:30)", "Functional",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "1", "phuong_thuc": "QR", "thoi_gian_quet": "08:45:00"}',
     "Status 200 OK, hanh_dong = 'CHECK_IN', trang_thai = 'DiMuon'",
     "Status 200 OK, hanh_dong = CHECK_IN, trang_thai = DiMuon", "PASSED", "Phân loại DiMuon chính xác"),

    (4, "TC-AttendanceScan", "TC_ATT_04", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Phân giải định danh dạng tiền tố INTERN-01", "Integration",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "INTERN-01"}',
     "Status 200 OK, tự động phân giải mã hồ sơ ma_ho_so = 1",
     "Status 200 OK, ma_ho_so = 1", "PASSED", "Bóc tách tiền tố chuỗi"),

    (5, "TC-AttendanceScan", "TC_ATT_05", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Phân giải định danh dạng tiền tố HS001", "Integration",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "HS001"}',
     "Status 200 OK, tự động phân giải mã hồ sơ ma_ho_so = 1",
     "Status 200 OK, ma_ho_so = 1", "PASSED", "Bóc tách tiền tố chuỗi"),

    (6, "TC-AttendanceScan", "TC_ATT_06", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Phân giải định danh trực tiếp bằng email sinh viên", "Integration",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "tranquochuy@ictu.edu.vn"}',
     "Status 200 OK, truy vấn hồ sơ qua email trong bảng NguoiDung",
     "Status 200 OK, tìm đúng ma_ho_so = 1", "PASSED", "Fallback email lookup"),

    (7, "TC-AttendanceScan", "TC_ATT_07", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Phân giải định danh bằng số điện thoại sinh viên", "Integration",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "0981234567"}',
     "Status 200 OK, truy vấn hồ sơ qua số điện thoại",
     "Status 200 OK, tìm đúng ma_ho_so = 1", "PASSED", "Fallback phone lookup"),

    (8, "TC-AttendanceScan", "TC_ATT_08", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Quét check-out về sớm (trước 17:00)", "Functional",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "1", "thoi_gian_quet": "16:30:00"}',
     "Status 200 OK, hanh_dong = 'CHECK_OUT', trang_thai = 'VeSom'",
     "Status 200 OK, hanh_dong = CHECK_OUT, trang_thai = VeSom", "PASSED", "Phân loại VeSom"),

    (9, "TC-AttendanceScan", "TC_ATT_09", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Quét check-out đúng giờ (sau 17:00)", "Functional",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "2", "thoi_gian_quet": "17:30:00"}',
     "Status 200 OK, hanh_dong = 'CHECK_OUT', trang_thai = 'DungGio', cập nhật gio_checkout",
     "Status 200 OK, CHECK_OUT thành công", "PASSED", "Hoàn tất ca làm việc"),

    (10, "TC-AttendanceScan", "TC_ATT_10", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Bản ghi chấm công lưu chuẩn xác vào bảng CSDL cham_cong", "Database",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "1", "phuong_thuc": "Card"}',
     "CSDL lưu đầy đủ ma_ho_so, ngay_cham_cong, gio_checkin, phuong_thuc, trang_thai",
     "Bản ghi cham_cong tồn tại trong CSDL với dữ liệu khớp", "PASSED", "Bảo toàn dữ liệu"),

    (11, "TC-AttendanceScan", "TC_ATT_11", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Bắt lỗi khi mã định danh rỗng hoặc toàn khoảng trắng", "Negative",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "   "}',
     "Status 422 Unprocessable Entity, báo lỗi mã định danh rỗng",
     "Status 422 Unprocessable Entity", "PASSED", "Validate Pydantic validator"),

    (12, "TC-AttendanceScan", "TC_ATT_12", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Bắt lỗi khi mã định danh không tồn tại trong hệ thống", "Negative",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "CARD-UNKNOWN-999999"}',
     "Status 404 Not Found, detail: 'Không tìm thấy hồ sơ thực tập sinh'",
     "Status 404 Not Found", "PASSED", "Bắt ngoại lệ 404"),

    (13, "TC-AttendanceScan", "TC_ATT_13", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Chặn chấm công thực tập sinh chưa được phê duyệt hồ sơ", "Business Rule",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "5"} (Hồ sơ ChoDuyet)',
     "Status 400 Bad Request, detail: 'chưa được phê duyệt'",
     "Status 400 Bad Request", "PASSED", "Quy tắc nghiệp vụ TTCS"),

    (14, "TC-AttendanceScan", "TC_ATT_14", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Chặn chấm công thực tập sinh đã thôi học (ThoiHoc)", "Business Rule",
     'POST /api/v1/attendance/scan {"ma_dinh_danh": "6"} (Trạng thái ThoiHoc)',
     "Status 400 Bad Request, detail: 'đã thôi học'",
     "Status 400 Bad Request", "PASSED", "Quy tắc nghiệp vụ TTCS"),

    (15, "TC-AttendanceScan", "TC_ATT_15", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Chống quét đúp liên tiếp ngay sau check-in (< 60 giây)", "Concurrency",
     'POST /api/v1/attendance/scan 2 lần cách nhau 10 giây',
     "Status 400 Bad Request, thông báo bạn vừa quét check-in, vui lòng không quét trùng",
     "Status 400 Bad Request, chặn quét đúp thành công", "PASSED", "Chống double-submission"),

    (16, "TC-AttendanceScan", "TC_ATT_16", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Chống quét lần 3 khi ca làm việc đã hoàn thành check-in và check-out", "Business Rule",
     'POST /api/v1/attendance/scan lần 3 trong ngày sau khi đã check-out',
     "Status 400 Bad Request, thông báo đã hoàn thành ca làm việc ngày hôm nay",
     "Status 400 Bad Request, chặn quét lần 3", "PASSED", "Khóa ca đã hoàn tất"),

    (17, "TC-AttendanceScan", "TC_ATT_17", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Không cho phép giờ check-out trước giờ check-in", "Time Logic",
     'Check-in 10:00:00, gửi quét check-out với thời gian 09:00:00',
     "Status 400 Bad Request, thời gian check-out không thể trước check-in",
     "Status 400 Bad Request", "PASSED", "Chặn thời gian nghịch lý"),

    (18, "TC-AttendanceScan", "TC_ATT_18", "Là thực tập sinh, quét mã QR / thẻ chấm công",
     "Thiếu trường bắt buộc ma_dinh_danh trong payload", "Schema Validation",
     'POST /api/v1/attendance/scan {"phuong_thuc": "QR"}',
     "Status 422 Unprocessable Entity, Pydantic báo lỗi thiếu field ma_dinh_danh",
     "Status 422 Unprocessable Entity", "PASSED", "Schema validation"),

    # --- NHIỆM VỤ 2: LÊN LỊCH HỌP & GỬI EMAIL HTML THÔNG BÁO TỰ ĐỘNG (18 TCs) ---
    (19, "TC-Meetings", "TC_MTG_01", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Tạo lịch họp đầy đủ thông tin hợp lệ", "Positive",
     'POST /api/v1/meetings (tieu_de, start, end, dia_diem, ma_nguoi_chu_tri, interns)',
     "Status 201 Created, code=201, trả về data lịch họp chi tiết",
     "Status 201 Created, data lịch họp đầy đủ", "PASSED", "Tạo lịch họp thành công"),

    (20, "TC-Meetings", "TC_MTG_02", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Tự động kích hoạt BackgroundTasks gửi email HTML thông báo tới sinh viên", "Background Tasks",
     'POST /api/v1/meetings mời thực tập sinh mã 3 (tranquochuy@ictu.edu.vn)',
     "BackgroundTasks gửi email không chặn API, sent_emails_history có email người nhận",
     "sent_emails_history chứa email gửi đến tranquochuy@ictu.edu.vn", "PASSED", "Tích hợp BackgroundTasks"),

    (21, "TC-Meetings", "TC_MTG_03", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Nội dung email text fallback và HTML chứa tiêu đề cuộc họp", "Email Content",
     'POST /api/v1/meetings {"tieu_de": "Họp giao ban đầu tuần"}',
     "Subject và nội dung email HTML chứa đúng tiêu đề cuộc họp",
     "Subject và HTML content chứa 'Họp giao ban đầu tuần'", "PASSED", "Mẫu email HTML chuyên nghiệp"),

    (22, "TC-Meetings", "TC_MTG_04", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Email chứa thông tin địa điểm và tên người chủ trì", "Email Content",
     'POST /api/v1/meetings {"dia_diem": "Phòng Lab AI", "ma_nguoi_chu_tri": 2}',
     "Email chứa địa điểm 'Phòng Lab AI' và tên người chủ trì 'Lê Quang Hưng'",
     "Nội dung email chứa đầy đủ địa điểm và người chủ trì", "PASSED", "Thông tin rõ ràng"),

    (23, "TC-Meetings", "TC_MTG_05", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Email HTML chứa thời gian bắt đầu và kết thúc định dạng chuẩn tiếng Việt", "Email Format",
     'POST /api/v1/meetings (ISO-8601 start/end)',
     "Thời gian trong email được format dạng '%H:%M ngày %d/%m/%Y' tiếng Việt",
     "Thời gian hiển thị tiếng Việt thân thiện", "PASSED", "Định dạng thời gian chuẩn"),

    (24, "TC-Meetings", "TC_MTG_06", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Tạo lịch họp trực tuyến kèm link Google Meet/Zoom trong địa điểm", "Functional",
     'POST /api/v1/meetings {"dia_diem": "https://meet.google.com/abc-defg-hij"}',
     "Status 201 Created, dia_diem lưu và gửi email kèm liên kết trực tuyến",
     "Status 201 Created, URL Google Meet hợp lệ", "PASSED", "Hỗ trợ họp Online"),

    (25, "TC-Meetings", "TC_MTG_07", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Hỗ trợ tiêu đề và nội dung cuộc họp bằng tiếng Việt có dấu đầy đủ UTF-8", "Encoding",
     'POST /api/v1/meetings {"tieu_de": "Hội thảo: Tối ưu hóa hiệu năng ứng dụng & Thiết kế hệ thống"}',
     "Status 201 Created, ký tự Unicode tiếng Việt được bảo toàn toàn vẹn",
     "Status 201 Created, hiển thị Unicode UTF-8 chuẩn xác", "PASSED", "Đa ngữ đi kèm"),

    (26, "TC-Meetings", "TC_MTG_08", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Thông tin cuộc họp được lưu chính xác trong bảng CSDL lich_hop", "Database",
     'POST /api/v1/meetings',
     "CSDL bảng lich_hop lưu đúng tieu_de, thoi_gian, dia_diem, ma_nguoi_chu_tri",
     "Bản ghi LichHop tồn tại trong CSDL với dữ liệu khớp", "PASSED", "Bảo toàn dữ liệu"),

    (27, "TC-Meetings", "TC_MTG_09", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Tự động loại bỏ ID trùng lặp trong danh_sach_thuc_tap_sinh", "De-duplication",
     'POST /api/v1/meetings {"danh_sach_thuc_tap_sinh": [3, 3, 3]} (Truyền trùng 3 lần)',
     "Status 201 Created, so_luong_nguoi_moi = 1, chỉ gửi 1 email duy nhất",
     "Status 201 Created, lọc trùng lặp thành công", "PASSED", "De-duplication"),

    (28, "TC-Meetings", "TC_MTG_10", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi khi tiêu đề cuộc họp để trống hoặc chỉ có khoảng trắng", "Negative",
     'POST /api/v1/meetings {"tieu_de": "   "}',
     "Status 400 hoặc 422, thông báo tiêu đề cuộc họp không được để trống",
     "Status 422 Unprocessable Entity", "PASSED", "Validate whitespace"),

    (29, "TC-Meetings", "TC_MTG_11", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi khi thiếu trường bắt buộc tieu_de trong payload", "Schema Validation",
     'POST /api/v1/meetings (Payload thiếu tieu_de)',
     "Status 422 Unprocessable Entity, Pydantic báo lỗi thiếu trường required",
     "Status 422 Unprocessable Entity", "PASSED", "Validate required fields"),

    (30, "TC-Meetings", "TC_MTG_12", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi thời gian kết thúc nhỏ hơn hoặc bằng thời gian bắt đầu", "Time Logic",
     'POST /api/v1/meetings (thoi_gian_ket_thuc < thoi_gian_bat_dau)',
     "Status 400 hoặc 422, thông báo thời gian kết thúc phải sau thời gian bắt đầu",
     "Status 422 Unprocessable Entity, chặn logic thời gian sai lệch", "PASSED", "Quy tắc GEMINI.md"),

    (31, "TC-Meetings", "TC_MTG_13", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi thời gian bắt đầu cuộc họp nằm trong quá khứ so với thời điểm tạo", "Business Rule",
     'POST /api/v1/meetings (thoi_gian_bat_dau = now - 2 ngày)',
     "Status 400 Bad Request, thời gian bắt đầu không được nằm trong quá khứ",
     "Status 400 Bad Request, chặn lịch họp quá khứ", "PASSED", "Quy tắc GEMINI.md"),

    (32, "TC-Meetings", "TC_MTG_14", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi người chủ trì không tồn tại trong CSDL", "Negative",
     'POST /api/v1/meetings {"ma_nguoi_chu_tri": 99999}',
     "Status 404 Not Found, detail: 'Không tìm thấy người chủ trì'",
     "Status 404 Not Found", "PASSED", "Kiểm tra khóa ngoại"),

    (33, "TC-Meetings", "TC_MTG_15", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi tài khoản người chủ trì đang bị khóa ('Khoa')", "Business Rule",
     'POST /api/v1/meetings với Host có trang_thai = "Khoa"',
     "Status 400 Bad Request, tài khoản người chủ trì bị khóa",
     "Status 400 Bad Request", "PASSED", "Kiểm soát trạng thái tài khoản"),

    (34, "TC-Meetings", "TC_MTG_16", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi khi danh sách thực tập sinh được mời rỗng []", "Validation",
     'POST /api/v1/meetings {"danh_sach_thuc_tap_sinh": []}',
     "Status 400 hoặc 422, danh sách sinh viên được mời không được để trống",
     "Status 422 Unprocessable Entity", "PASSED", "Validate min_items >= 1"),

    (35, "TC-Meetings", "TC_MTG_17", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi danh sách sinh viên được mời chứa mã không tồn tại trong CSDL", "Negative",
     'POST /api/v1/meetings {"danh_sach_thuc_tap_sinh": [3, 88888]}',
     "Status 404 Not Found, không tìm thấy sinh viên mã 88888",
     "Status 404 Not Found", "PASSED", "Bắt trọn lỗi sinh viên ảo"),

    (36, "TC-Meetings", "TC_MTG_18", "Là Mentor/HR, tạo lịch họp và gửi email thông báo",
     "Bắt lỗi bỏ trống địa điểm cuộc họp", "Validation",
     'POST /api/v1/meetings {"dia_diem": "   "}',
     "Status 400 hoặc 422, địa điểm không được để trống hoặc chỉ có khoảng trắng",
     "Status 422 Unprocessable Entity", "PASSED", "Validate địa điểm"),

    # --- NHIỆM VỤ 3: SAO LƯU DỮ LIỆU HỆ THỐNG TỰ ĐỘNG (18 TCs) ---
    (37, "TC-SystemBackup", "TC_BK_01", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Kiểm tra hàm format bytes chuyển đổi dung lượng thân thiện", "Utility Unit",
     "format_bytes(500), format_bytes(2048), format_bytes(5MB), format_bytes(2GB)",
     "Chuyển đổi đúng các đơn vị '500 B', '2.00 KB', '5.00 MB', '2.00 GB'",
     "Định dạng chuỗi dung lượng chuẩn xác", "PASSED", "Helper format bytes"),

    (38, "TC-SystemBackup", "TC_BK_02", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Kiểm tra xác thực tên file sao lưu hợp lệ có đuôi .sql.gz", "Validation",
     "validate_backup_filename('backup_20261019.sql.gz')",
     "Trả về tên file an toàn nguyên vẹn",
     "Tên file hợp lệ được chấp nhận", "PASSED", "Whitelist regex filename"),

    (39, "TC-SystemBackup", "TC_BK_03", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Chặn tên file sao lưu chứa ký tự đường dẫn Path Traversal (../ hoặc ..\\)", "Security",
     "validate_backup_filename('../../../etc/passwd.sql.gz')",
     "Ném BackupSecurityError, chặn đứng tấn công đọc ghi file trái phép",
     "Ném BackupSecurityError thành công", "PASSED", "Bảo mật Path Traversal"),

    (40, "TC-SystemBackup", "TC_BK_04", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Chặn tên file sao lưu có phần mở rộng độc hại (.exe, .bat, .sh)", "Security",
     "validate_backup_filename('hack.exe'), 'script.sh', 'test.sql'",
     "Ném BackupSecurityError, chỉ chấp nhận đuôi .sql.gz duy nhất",
     "Ném BackupSecurityError cho mọi extension lạ", "PASSED", "Chặn tệp thực thi độc hại"),

    (41, "TC-SystemBackup", "TC_BK_05", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Tính toán mã băm SHA-256 chính xác 64 ký tự hex cho file sao lưu", "Data Integrity",
     "calculate_sha256(test_file)",
     "Trả về chuỗi hash SHA-256 đúng 64 ký tự hex hợp lệ",
     "Chuỗi hex 64 ký tự chính xác", "PASSED", "Toàn vẹn Checksum SHA-256"),

    (42, "TC-SystemBackup", "TC_BK_06", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Admin kích hoạt tạo bản sao lưu thủ công thành công", "Positive",
     'POST /api/v1/system/backup (Header X-User-Email: admin@example.com)',
     "Status 201 Created, trả về filename, size_bytes, download_url",
     "Status 201 Created, tạo file backup thành công", "PASSED", "Tạo backup thủ công"),

    (43, "TC-SystemBackup", "TC_BK_07", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "File sao lưu tạo ra được nén bằng chuẩn gzip hợp lệ và giải nén an toàn", "File System",
     "POST /api/v1/system/backup -> Đọc gzip.open(filepath)",
     "File tồn tại trên đĩa, đọc và giải nén dữ liệu gzip thành công",
     "File .sql.gz hợp lệ, cấu trúc nén nguyên vẹn", "PASSED", "Chuẩn nén gzip"),

    (44, "TC-SystemBackup", "TC_BK_08", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Bản sao lưu ghi nhận đầy đủ mã băm SHA-256 và dung lượng thực tế", "Data Integrity",
     "POST /api/v1/system/backup",
     "Metadata có checksum_sha256 (64 hex), size_bytes > 0, size_human",
     "Metadata đầy đủ hash và dung lượng thực tế", "PASSED", "Metadata toàn vẹn"),

    (45, "TC-SystemBackup", "TC_BK_09", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Lưu trữ chính xác ghi chú note đính kèm khi tạo backup", "Functional",
     'POST /api/v1/system/backup {"note": "Sao lưu tự động định kỳ cuối ngày Sprint 4"}',
     "Response và metadata lưu đúng nội dung ghi chú custom note",
     "Ghi chú note được lưu nguyên vẹn", "PASSED", "Custom note"),

    (46, "TC-SystemBackup", "TC_BK_10", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Lấy danh sách các bản sao lưu hiện có -> HTTP 200 OK", "Functional",
     "GET /api/v1/system/backups (Header Admin)",
     "Status 200 OK, trả về danh sách backups mảng và total_files",
     "Status 200 OK, danh sách tệp sao lưu", "PASSED", "Truy vấn danh sách backup"),

    (47, "TC-SystemBackup", "TC_BK_11", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Danh sách sao lưu được sắp xếp theo thời gian mới nhất lên đầu", "Functional",
     "GET /api/v1/system/backups",
     "Tệp tạo sau có created_at >= tệp tạo trước (sắp xếp DESC)",
     "Danh sách sắp xếp giảm dần chuẩn xác", "PASSED", "Sắp xếp DESC"),

    (48, "TC-SystemBackup", "TC_BK_12", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Header Content-Disposition chứa attachment và tên file khi tải về", "Streaming",
     "GET /api/v1/system/backups/{filename}/download",
     "Status 200 OK, Content-Disposition có 'attachment; filename=...'",
     "Header tải về đính kèm tệp chuẩn", "PASSED", "Stream chunk 64KB"),

    (49, "TC-SystemBackup", "TC_BK_13", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Xóa bản sao lưu thành công (DELETE) và kiểm tra file bị xóa", "Functional",
     "DELETE /api/v1/system/backups/{filename}",
     "Status 200 OK, file vật lý bị xóa khỏi thư mục (tải lại ra 404)",
     "Status 200 OK, file đã bị xóa", "PASSED", "Xóa bản backup cũ"),

    (50, "TC-SystemBackup", "TC_BK_14", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Chặn tải file chứa ký tự đường dẫn ../ (Path Traversal Security)", "Security",
     "GET /api/v1/system/backups/..%2F..%2Fmain.py/download",
     "Status 400 hoặc 403 hoặc 404, tuyệt đối không lộ file mã nguồn hệ thống",
     "Status 400/404, chặn Path Traversal", "PASSED", "An toàn bảo mật"),

    (51, "TC-SystemBackup", "TC_BK_15", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Tải file sao lưu không tồn tại trong thư mục", "Negative",
     "GET /api/v1/system/backups/backup_fake_99999.sql.gz/download",
     "Status 404 Not Found, detail: 'Không tìm thấy tệp sao lưu'",
     "Status 404 Not Found", "PASSED", "Xử lý 404 file ảo"),

    (52, "TC-SystemBackup", "TC_BK_16", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Xóa file sao lưu không tồn tại", "Negative",
     "DELETE /api/v1/system/backups/backup_fake_file_99999.sql.gz",
     "Status 404 Not Found, không tìm thấy file để xóa",
     "Status 404 Not Found", "PASSED", "Xử lý 404 file ảo"),

    (53, "TC-SystemBackup", "TC_BK_17", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Phân quyền RBAC: Chặn tài khoản Thực tập sinh gọi API sao lưu", "Security RBAC",
     "POST /api/v1/system/backup (Header TTS: tranquochuy@ictu.edu.vn)",
     "Status 403 Forbidden, không đủ quyền quản trị hệ thống",
     "Status 403 Forbidden", "PASSED", "Chặn non-admin"),

    (54, "TC-SystemBackup", "TC_BK_18", "Là Admin, sao lưu dữ liệu tự động định kỳ",
     "Chống xung đột đồng thời: Chặn tạo backup khi tiến trình khác đang chạy", "Concurrency Lock",
     "POST /api/v1/system/backup khi _is_backup_running = True",
     "Status 409 Conflict, thông báo tiến trình sao lưu khác đang được thực hiện",
     "Status 409 Conflict, Concurrency Lock bảo vệ an toàn", "PASSED", "Concurrency Lock 409")
]

for r_idx, row in enumerate(tc_54_data, start=2):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate(row, start=1):
        cell = ws2.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = border_cell
        if fill.fill_type: cell.fill = fill
        
        # Căn lề
        if c_idx in (1, 2, 3, 6, 10):
            cell.alignment = align_center
        else:
            cell.alignment = align_wrap_left

        # Highlight cột PASSED
        if c_idx == 10 and val == "PASSED":
            cell.font = font_pass
            cell.fill = fill_pass

# Độ rộng cột Sheet 2
col_widths_s2 = {
    'A': 8.0, 'B': 18.0, 'C': 15.0, 'D': 36.0, 'E': 46.0,
    'F': 16.0, 'G': 46.0, 'H': 46.0, 'I': 46.0, 'J': 14.0, 'K': 18.0
}
for col_letter, width in col_widths_s2.items():
    ws2.column_dimensions[col_letter].width = width


# ==============================================================================
# SHEET 3: BÁO CÁO LỖI & ĐỀ XUẤT
# ==============================================================================
ws3 = wb.create_sheet(title="Báo Cáo Lỗi & Đề Xuất")
ws3.views.sheetView[0].showGridLines = True

note_headers = [
    "Mã Ghi Nhận", "Test Key", "Tiêu đề / Nội dung ghi nhận (Summary)",
    "Mức độ (Severity)", "Độ ưu tiên (Priority)", "Phân loại",
    "Các bước kiểm tra / Tái hiện (Steps)", "Kết quả thực tế (Actual Result)",
    "Kết quả kỳ vọng / Đề xuất tối ưu (Expected)", "Người báo cáo",
    "Phụ trách khắc phục", "Trạng thái (Status)"
]

for c_idx, h in enumerate(note_headers, start=1):
    cell = ws3.cell(1, c_idx, h)
    cell.font = font_header
    cell.fill = fill_navy_header
    cell.alignment = align_center
    cell.border = border_header

notes_data = [
    ("NOTE-01", "TC-AttendanceScan",
     "Kiểm soát chặt chẽ cơ chế chống quét trùng lặp trong 60 giây và khóa ca hoàn thành",
     "Medium", "High", "Kiểm soát nghiệp vụ & Concurrency",
     "1. Quét check-in lúc 08:15:00\n2. Gửi tiếp quét check-in lúc 08:15:10 (< 60s)\n3. Quét check-out lúc 17:30:00\n4. Quét lần 3 trong ngày",
     "Hệ thống trả về HTTP 400 Bad Request kèm thông điệp rõ ràng chặn quét đúp và chặn quét lần 3 khi ca đã hoàn thành",
     "Đạt chuẩn nghiệp vụ TTCS và GEMINI.md: chống double submission và kiểm soát chặt chẽ trạng thái chấm công",
     "Trần Quốc Huy", "Nguyễn Văn Hiếu (Backend)", "Verified / Đã xác nhận"),

    ("NOTE-02", "TC-Meetings",
     "Bắt trọn vẹn ngoại lệ thời gian kết thúc trước thời gian bắt đầu và thời gian trong quá khứ",
     "Medium", "Medium", "Quy tắc nghiệp vụ thời gian",
     "1. Gọi POST /api/v1/meetings với thoi_gian_ket_thuc <= thoi_gian_bat_dau\n2. Gọi POST /api/v1/meetings với thoi_gian_bat_dau trong quá khứ",
     "Backend và Pydantic validator chặn đứng và trả về HTTP 400/422 với thông báo lỗi tiếng Việt dễ hiểu",
     "Đúng quy định GEMINI.md Mục 1.A.2: thoi_gian_ket_thuc phải lớn hơn thoi_gian_bat_dau và không được nằm trong quá khứ",
     "Trần Quốc Huy", "Nguyễn Văn Hiếu (Backend)", "Verified / Đã xác nhận"),

    ("NOTE-03", "TC-SystemBackup",
     "Bảo mật tuyệt đối chống tấn công Path Traversal (../) và khóa chống xung đột sao lưu Concurrency Lock (409)",
     "High", "High", "Bảo mật hệ thống & Concurrency",
     "1. Thử tải file sao lưu chứa ký tự đường dẫn: ..%2F..%2Fmain.py\n2. Gọi kích hoạt sao lưu khi một tiến trình sao lưu khác đang chạy",
     "Hệ thống chặn đứng Path Traversal (HTTP 400/404) và ném HTTP 409 Conflict khi phát hiện Concurrency Lock đang active",
     "Đảm bảo an toàn tuyệt đối cho tệp mã nguồn và CSDL hệ thống theo tiêu chuẩn bảo mật doanh nghiệp",
     "Trần Quốc Huy", "Nguyễn Văn Hiếu (Backend)", "Verified / Đã xác nhận")
]

for r_idx, row in enumerate(notes_data, start=2):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate(row, start=1):
        cell = ws3.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = border_cell
        if fill.fill_type: cell.fill = fill
        if c_idx in (1, 2, 4, 5, 6, 10, 11, 12):
            cell.alignment = align_center
        else:
            cell.alignment = align_wrap_left
        if c_idx == 12:
            cell.font = font_pass
            cell.fill = fill_pass

# Độ rộng cột Sheet 3
col_widths_s3 = {
    'A': 12.0, 'B': 18.0, 'C': 35.0, 'D': 14.0, 'E': 14.0, 'F': 22.0,
    'G': 35.0, 'H': 35.0, 'I': 35.0, 'J': 16.0, 'K': 22.0, 'L': 20.0
}
for col_letter, width in col_widths_s3.items():
    ws3.column_dimensions[col_letter].width = width

# Lưu file kết quả
out_filename = "Bao_Cao_Kiem_Thu_Sprint4_Tran_Quoc_Huy.xlsx"
wb.save(out_filename)
import sys
try:
    print(f"Bao cao Excel da duoc tao thanh cong: {out_filename}")
except Exception:
    pass
