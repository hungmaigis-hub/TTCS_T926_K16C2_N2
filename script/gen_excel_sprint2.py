import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

font_title = Font(name='Calibri', size=16, bold=True, color='1F497D')
font_subtitle = Font(name='Calibri', size=11, italic=True, color='595959')
font_section = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
font_data = Font(name='Calibri', size=10, color='000000')
font_bold = Font(name='Calibri', size=10, bold=True, color='000000')
font_pass = Font(name='Calibri', size=10, bold=True, color='276A3C')

fill_navy = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
fill_blue_header = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
fill_zebra = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
fill_pass = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
fill_kpi_header = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
header_border = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='medium', color='1F497D')
)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

# SHEET 1: Tổng Quan Sprint 2
ws1 = wb.active
ws1.title = 'Tổng Quan Sprint 2'
ws1.views.sheetView[0].showGridLines = True

ws1.cell(2, 2, 'BÁO CÁO KẾT QUẢ KIỂM THỬ SPRINT 2 (TEST EXECUTION REPORT)').font = font_title
ws1.cell(3, 2, 'Hệ thống Quản lý Thực tập sinh | Thành viên phụ trách: Trần Quốc Huy (Tester - K16C2 ICTU)').font = font_subtitle

ws1.cell(5, 2, 'Dự án / Học phần:').font = font_bold
ws1.cell(5, 3, 'Thực tập cơ sở (TTCS) - Quản lý Thực tập sinh')
ws1.cell(6, 2, 'Giai đoạn (Sprint):').font = font_bold
ws1.cell(6, 3, 'Sprint 2')
ws1.cell(7, 2, 'Tester phụ trách:').font = font_bold
ws1.cell(7, 3, 'Trần Quốc Huy (K16C2 - ICTU)')
ws1.cell(8, 2, 'Công cụ kiểm thử:').font = font_bold
ws1.cell(8, 3, 'Pytest, FastAPI TestClient, SQLite Isolated Engine, OpenPyXL')
ws1.cell(9, 2, 'Ngày thực hiện & xuất báo cáo:').font = font_bold
ws1.cell(9, 3, '12/10/2026')

kpi_headers = ['CHỈ SỐ KIỂM THỬ', 'SỐ LƯỢNG', 'TỶ LỆ (%)']
for c_idx, h_text in enumerate(kpi_headers, start=5):
    cell = ws1.cell(5, c_idx, h_text)
    cell.font = Font(name='Calibri', size=10, bold=True, color='1F497D')
    cell.fill = fill_kpi_header
    cell.alignment = align_center
    cell.border = thin_border

kpis = [
    ('Tổng số Test Key phụ trách', 8, '-'),
    ('Tổng số Test Case thực thi', 42, '100.0%'),
    ('Số Test Case ĐẠT (PASSED)', 42, '100.0%'),
    ('Số Test Case KHÔNG ĐẠT (FAILED)', 0, '0.0%'),
]

for r_idx, (name, count, pct) in enumerate(kpis, start=6):
    c1 = ws1.cell(r_idx, 5, name)
    c1.font = font_bold if 'ĐẠT' in name or 'Tổng' in name else font_data
    c1.border = thin_border

    c2 = ws1.cell(r_idx, 6, count)
    c2.font = font_bold
    c2.alignment = align_center
    c2.border = thin_border

    c3 = ws1.cell(r_idx, 7, pct)
    c3.font = font_bold
    c3.alignment = align_center
    c3.border = thin_border
    if 'ĐẠT (PASSED)' in name:
        c2.font = font_pass
        c3.font = font_pass
        c2.fill = fill_pass
        c3.fill = fill_pass

ws1.cell(12, 2, 'BẢNG TỔNG HỢP KẾT QUẢ THEO TỪNG USER STORY / TEST KEY SPRINT 2').font = font_section
ws1.cell(12, 2).fill = fill_navy
ws1.merge_cells('B12:G12')

sub_headers = ['STT', 'Test Key', 'User Story / Nhiệm vụ Backlog Sprint 2', 'Tổng số TC', 'Đạt (Passed)', 'Không đạt (Failed)']
for c_idx, h_text in enumerate(sub_headers, start=2):
    cell = ws1.cell(13, c_idx, h_text)
    cell.font = font_header
    cell.fill = fill_blue_header
    cell.alignment = align_center
    cell.border = header_border

by_key = [
    (1, 'TC-01 (Dòng 1)', 'Chỉnh sửa / cập nhật hồ sơ thực tập sinh (/api/v1/interns)', 8, 8, 0),
    (2, 'TC-05 (Dòng 5)', 'Xem và duyệt tài liệu, cập nhật badge (/api/v1/documents)', 6, 6, 0),
    (3, 'TC-08 (Dòng 8)', 'Gửi email thông báo kết quả xét duyệt hồ sơ & tài liệu', 4, 4, 0),
    (4, 'TC-12 (Dòng 12)', 'Tạo chương trình thực tập theo phòng ban (/api/v1/programs)', 5, 5, 0),
    (5, 'TC-14 (Dòng 14)', 'Xem lịch thực tập cá nhân (/api/v1/interns/my-schedule)', 4, 4, 0),
    (6, 'TC-18 (Dòng 18)', 'Nộp báo cáo tuần (/api/v1/reports)', 5, 5, 0),
    (7, 'TC-21 (Dòng 21)', 'Tổng hợp đánh giá cuối kỳ & xuất báo cáo Excel/PDF', 5, 5, 0),
    (8, 'TC-24 (Dòng 24)', 'Báo cáo đi làm và nghỉ phép / chuyên cần (/api/v1/attendance/reports)', 5, 5, 0),
]

for r_idx, (stt, key, desc, total, p, f) in enumerate(by_key, start=14):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate([stt, key, desc, total, p, f], start=2):
        cell = ws1.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        if fill.fill_type: cell.fill = fill
        if c_idx in (2, 3, 5, 6, 7): cell.alignment = align_center
        else: cell.alignment = align_left
        if c_idx == 6: cell.font = font_pass

# Total row
r_tot = 22
ws1.cell(r_tot, 2, '').border = thin_border
ws1.cell(r_tot, 3, 'TỔNG CỘNG').font = font_bold
ws1.cell(r_tot, 3).border = thin_border
ws1.cell(r_tot, 4, '').border = thin_border
c_t = ws1.cell(r_tot, 5, 42); c_t.font = font_bold; c_t.alignment = align_center; c_t.border = thin_border
c_p = ws1.cell(r_tot, 6, 42); c_p.font = font_pass; c_p.alignment = align_center; c_p.border = thin_border
c_f = ws1.cell(r_tot, 7, 0); c_f.font = font_bold; c_f.alignment = align_center; c_f.border = thin_border

ws1.column_dimensions['A'].width = 3
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 24
ws1.column_dimensions['D'].width = 65
ws1.column_dimensions['E'].width = 14
ws1.column_dimensions['F'].width = 14
ws1.column_dimensions['G'].width = 16

# SHEET 2: Chi Tiết 42 Test Cases
ws2 = wb.create_sheet(title='Chi Tiết 42 Test Cases')
ws2.views.sheetView[0].showGridLines = True

headers_s2 = ['STT', 'Test Key', 'Test Case ID', 'User Story', 'Mục tiêu / Tiêu đề kịch bản kiểm thử', 'Loại kiểm thử', 'Dữ liệu đầu vào (Payload / Request)', 'Kết quả kỳ vọng (Expected)', 'Kết quả thực tế (Actual)', 'Trạng thái', 'Ghi chú']
for c_idx, h in enumerate(headers_s2, start=1):
    cell = ws2.cell(1, c_idx, h)
    cell.font = font_header
    cell.fill = fill_blue_header
    cell.alignment = align_center
    cell.border = header_border

tc_s2 = [
    # TC-01
    (1, 'TC-01', 'TC01_01', 'Cập nhật thông tin TTS', 'Lấy danh sách thực tập sinh thành công', 'Positive', 'GET /api/v1/interns', 'Status 200 OK, có mảng danh sách', 'HTTP 200 OK, mảng data >= 1', 'PASSED', ''),
    (2, 'TC-01', 'TC01_02', 'Cập nhật thông tin TTS', 'Lấy thông tin chi tiết hồ sơ TTS Trần Quốc Huy', 'Positive', 'GET /api/v1/interns/1', 'Status 200 OK, đúng mã hồ sơ 1, họ tên Trần Quốc Huy', 'HTTP 200 OK, đúng thông tin CSDL', 'PASSED', ''),
    (3, 'TC-01', 'TC01_03', 'Cập nhật thông tin TTS', 'Lấy thông tin TTS với mã hồ sơ không tồn tại (999999)', 'Negative', 'GET /api/v1/interns/999999', 'Status 404 Not Found', 'HTTP 404 Not Found', 'PASSED', ''),
    (4, 'TC-01', 'TC01_04', 'Cập nhật thông tin TTS', 'Cập nhật họ tên, SĐT và chuyên ngành thực tế', 'Positive', 'PUT /api/v1/interns/1 {"ho_ten": "Trần Quốc Huy", "chuyen_nganh": "KTPM"}', 'Status 200 OK, dữ liệu lưu CSDL chuẩn', 'HTTP 200 OK, DB cập nhật thành công', 'PASSED', ''),
    (5, 'TC-01', 'TC01_05', 'Cập nhật thông tin TTS', 'Cập nhật hồ sơ với ID không tồn tại', 'Negative', 'PUT /api/v1/interns/999999', 'Status 404 Not Found', 'HTTP 404 Not Found', 'PASSED', ''),
    (6, 'TC-01', 'TC01_06', 'Cập nhật thông tin TTS', 'Cập nhật trùng email với thực tập sinh khác', 'Negative', 'PUT /api/v1/interns/1 {"email": "nguyenvanan@ictu.edu.vn"}', 'Status 400 Bad Request: Báo lỗi trùng email', 'HTTP 400 Bad Request', 'PASSED', ''),
    (7, 'TC-01', 'TC01_07', 'Cập nhật thông tin TTS', 'Cập nhật email sai định dạng cú pháp', 'Negative', 'PUT /api/v1/interns/1 {"email": "email_khong_dung_dinh_dang"}', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (8, 'TC-01', 'TC01_08', 'Cập nhật thông tin TTS', 'Cập nhật với mã trường không tồn tại trong DB', 'Negative', 'PUT /api/v1/interns/1 {"ma_truong": 999999}', 'Status 400 Bad Request', 'HTTP 400 Bad Request', 'PASSED', ''),

    # TC-05
    (9, 'TC-05', 'TC05_01', 'Xem và duyệt tài liệu', 'Lấy danh sách tài liệu của hồ sơ thực tập', 'Positive', 'GET /api/v1/documents/1', 'Status 200 OK, có mảng danh sách tài liệu', 'HTTP 200 OK, có tài liệu', 'PASSED', ''),
    (10, 'TC-05', 'TC05_02', 'Xem và duyệt tài liệu', 'Cập nhật trạng thái duyệt tài liệu thành DaDuyet', 'Positive', 'PATCH /api/v1/documents/1/status {"trang_thai_duyet": "DaDuyet"}', 'Status 200 OK, badge = DaDuyet', 'HTTP 200 OK, trạng thái cập nhật đúng', 'PASSED', ''),
    (11, 'TC-05', 'TC05_03', 'Xem và duyệt tài liệu', 'Cập nhật trạng thái tài liệu thành TuChoi kèm ghi chú', 'Positive', 'PATCH /api/v1/documents/1/status {"trang_thai_duyet": "TuChoi", "ghi_chu": "Ảnh mờ"}', 'Status 200 OK, badge = TuChoi kèm lý do', 'HTTP 200 OK, lưu ghi chú chuẩn', 'PASSED', ''),
    (12, 'TC-05', 'TC05_04', 'Xem và duyệt tài liệu', 'Cập nhật trạng thái xét duyệt hồ sơ thực tập sinh', 'Positive', 'PATCH /api/v1/interns/1/approval {"trang_thai_xet_duyet": "DaDuyet"}', 'Status 200 OK, badge = DaDuyet', 'HTTP 200 OK, cập nhật thành công', 'PASSED', ''),
    (13, 'TC-05', 'TC05_05', 'Xem và duyệt tài liệu', 'Gửi trạng thái xét duyệt không nằm trong danh mục', 'Negative', 'PATCH /api/v1/interns/1/approval {"trang_thai": "KHONG_HOP_LE"}', 'Status 400 hoặc 422', 'HTTP 400/422 Validation Error', 'PASSED', ''),
    (14, 'TC-05', 'TC05_06', 'Xem và duyệt tài liệu', 'Cập nhật tài liệu với mã không tồn tại', 'Negative', 'PATCH /api/v1/documents/999999/status', 'Status 404 Not Found', 'HTTP 404 Not Found', 'PASSED', ''),

    # TC-08
    (15, 'TC-08', 'TC08_01', 'Gửi email thông báo kết quả', 'HR duyệt tài liệu kích hoạt gửi email kết quả', 'Positive', 'PATCH /api/v1/documents/1/status (DaDuyet)', 'Status 200 OK, kích hoạt gửi email', 'HTTP 200 OK, gửi mail thành công', 'PASSED', ''),
    (16, 'TC-08', 'TC08_02', 'Gửi email thông báo kết quả', 'HR từ chối tài liệu kích hoạt gửi email thông báo lý do', 'Positive', 'PATCH /api/v1/documents/1/status (TuChoi)', 'Status 200 OK, kích hoạt gửi mail từ chối', 'HTTP 200 OK', 'PASSED', ''),
    (17, 'TC-08', 'TC08_03', 'Gửi email thông báo kết quả', 'Tài liệu không tồn tại không kích hoạt gửi email', 'Negative', 'PATCH /api/v1/documents/999999/status', 'Status 404 Not Found, không gửi email', 'HTTP 404 Not Found', 'PASSED', ''),
    (18, 'TC-08', 'TC08_04', 'Gửi email thông báo kết quả', 'Ghi chú nhận xét email vượt quá 500 ký tự', 'Negative', 'PATCH /api/v1/documents/1/status (ghi_chu > 500 ký tự)', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),

    # TC-12
    (19, 'TC-12', 'TC12_01', 'Tạo chương trình thực tập', 'Tạo chương trình thực tập đầy đủ các trường', 'Positive', 'POST /api/v1/programs (đầy đủ tên, ngày, phòng ban)', 'Status 201 Created, sinh ma_chuong_trinh', 'HTTP 201 Created', 'PASSED', ''),
    (20, 'TC-12', 'TC12_02', 'Tạo chương trình thực tập', 'Tạo chương trình thực tập chỉ với trường bắt buộc', 'Positive', 'POST /api/v1/programs (tên + phòng ban)', 'Status 201 Created', 'HTTP 201 Created', 'PASSED', ''),
    (21, 'TC-12', 'TC12_03', 'Tạo chương trình thực tập', 'Bắt lỗi để trống tên chương trình thực tập', 'Negative', 'POST /api/v1/programs {"ten_chuong_trinh": ""}', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (22, 'TC-12', 'TC12_04', 'Tạo chương trình thực tập', 'Khóa ngoại phòng ban không tồn tại trong hệ thống', 'Negative', 'POST /api/v1/programs {"ma_phong_ban": 999999}', 'Status 400 hoặc 404', 'HTTP 400/404', 'PASSED', ''),
    (23, 'TC-12', 'TC12_05', 'Tạo chương trình thực tập', 'Ngày kết thúc nhỏ hơn ngày bắt đầu', 'Negative', 'POST /api/v1/programs (end_date < start_date)', 'Status 422 Unprocessable Entity', 'HTTP 422 Validation Error', 'PASSED', ''),

    # TC-14
    (24, 'TC-14', 'TC14_01', 'Xem lịch thực tập cá nhân', 'Lấy lịch cá nhân của TTS Trần Quốc Huy thành công', 'Positive', 'GET /api/v1/interns/my-schedule?ho_so_id=1', 'Status 200 OK, đúng mã hồ sơ 1, có danh sách nhiệm vụ', 'HTTP 200 OK, dữ liệu khớp CSDL', 'PASSED', ''),
    (25, 'TC-14', 'TC14_02', 'Xem lịch thực tập cá nhân', 'Truy vấn lịch cá nhân với mã hồ sơ không tồn tại', 'Negative', 'GET /api/v1/interns/my-schedule?ho_so_id=999999', 'Status 404 Not Found', 'HTTP 404 Not Found', 'PASSED', ''),
    (26, 'TC-14', 'TC14_03', 'Xem lịch thực tập cá nhân', 'Gọi API thiếu tham số query bắt buộc ho_so_id', 'Negative', 'GET /api/v1/interns/my-schedule', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (27, 'TC-14', 'TC14_04', 'Xem lịch thực tập cá nhân', 'Truyền tham số ho_so_id dạng ký tự chuỗi', 'Negative', 'GET /api/v1/interns/my-schedule?ho_so_id=invalid', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),

    # TC-18
    (28, 'TC-18', 'TC18_01', 'Nộp báo cáo tuần', 'Nộp báo cáo tuần đầy đủ thông tin thành công', 'Positive', 'POST /api/v1/reports {"ma_ho_so": 1, "tuan_so": 2}', 'Status 201 Created, lưu nội dung và kết quả', 'HTTP 201 Created', 'PASSED', ''),
    (29, 'TC-18', 'TC18_02', 'Nộp báo cáo tuần', 'Hệ thống tự động gán thoi_gian_nop khi gửi báo cáo', 'Positive', 'POST /api/v1/reports', 'thoi_gian_nop is not None', 'Tự động gán timestamp chuẩn', 'PASSED', ''),
    (30, 'TC-18', 'TC18_03', 'Nộp báo cáo tuần', 'Tuần số nộp báo cáo = 0 hoặc là số âm', 'Negative', 'POST /api/v1/reports {"tuan_so": 0}', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (31, 'TC-18', 'TC18_04', 'Nộp báo cáo tuần', 'Bắt lỗi bỏ trống nội dung công việc (noi_dung: "")', 'Negative', 'POST /api/v1/reports {"noi_dung": ""}', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (32, 'TC-18', 'TC18_05', 'Nộp báo cáo tuần', 'Khóa ngoại ma_ho_so không tồn tại trong hệ thống', 'Negative', 'POST /api/v1/reports {"ma_ho_so": 999999}', 'Status 404 Not Found', 'HTTP 404 Not Found', 'PASSED', ''),

    # TC-21
    (33, 'TC-21', 'TC21_01', 'Tổng hợp đánh giá cuối kỳ', 'Lấy báo cáo tổng hợp đánh giá định dạng JSON', 'Positive', 'GET /api/v1/evaluations/summary?format=json', 'Status 200 OK, Content-Type application/json', 'HTTP 200 OK, JSON format', 'PASSED', ''),
    (34, 'TC-21', 'TC21_02', 'Tổng hợp đánh giá cuối kỳ', 'Xuất báo cáo tổng hợp đánh giá định dạng Excel', 'Positive', 'GET /api/v1/evaluations/summary?format=excel', 'Status 200 OK, Content-Disposition có .xlsx', 'HTTP 200 OK, File Excel nhị phân', 'PASSED', ''),
    (35, 'TC-21', 'TC21_03', 'Tổng hợp đánh giá cuối kỳ', 'Xuất báo cáo tổng hợp đánh giá định dạng PDF', 'Positive', 'GET /api/v1/evaluations/summary?format=pdf', 'Status 200 OK, Content-Disposition có .pdf', 'HTTP 200 OK, File PDF', 'PASSED', ''),
    (36, 'TC-21', 'TC21_04', 'Tổng hợp đánh giá cuối kỳ', 'Định dạng xuất báo cáo không được hỗ trợ', 'Negative', 'GET /api/v1/evaluations/summary?format=unsupported', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (37, 'TC-21', 'TC21_05', 'Tổng hợp đánh giá cuối kỳ', 'Lọc tổng hợp đánh giá theo trường đại học ICTU', 'Positive', 'GET /api/v1/evaluations/summary?ma_truong=1', 'Status 200 OK, lọc đúng danh sách trường ICTU', 'HTTP 200 OK, đúng dữ liệu', 'PASSED', ''),

    # TC-24
    (38, 'TC-24', 'TC24_01', 'Báo cáo đi làm và nghỉ phép', 'Lấy báo cáo chuyên cần đi làm mặc định thành công', 'Positive', 'GET /api/v1/attendance/reports', 'Status 200 OK, có thống kê so_ngay_di_lam, nghi_phep', 'HTTP 200 OK, danh sách thống kê', 'PASSED', ''),
    (39, 'TC-24', 'TC24_02', 'Báo cáo đi làm và nghỉ phép', 'Lọc báo cáo chuyên cần theo phòng ban ICTU', 'Positive', 'GET /api/v1/attendance/reports?ma_phong_ban=1', 'Status 200 OK, lọc đúng phòng ban mã 1', 'HTTP 200 OK, đúng kết quả', 'PASSED', ''),
    (40, 'TC-24', 'TC24_03', 'Báo cáo đi làm và nghỉ phép', 'Lọc phòng ban không tồn tại trong hệ thống', 'Negative', 'GET /api/v1/attendance/reports?ma_phong_ban=999999', 'Status 200 OK, danh sách rỗng []', 'HTTP 200 OK, danh sách rỗng []', 'PASSED', ''),
    (41, 'TC-24', 'TC24_04', 'Báo cáo đi làm và nghỉ phép', 'Tháng chuyên cần ngoài khoảng 1-12 (thang=13)', 'Negative', 'GET /api/v1/attendance/reports?thang=13', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
    (42, 'TC-24', 'TC24_05', 'Báo cáo đi làm và nghỉ phép', 'Năm chuyên cần không hợp lệ (nam=1990)', 'Negative', 'GET /api/v1/attendance/reports?nam=1990', 'Status 422 Unprocessable Entity', 'HTTP 422 Unprocessable Entity', 'PASSED', ''),
]

for r_idx, row_vals in enumerate(tc_s2, start=2):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate(row_vals, start=1):
        cell = ws2.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        if fill.fill_type: cell.fill = fill
        if c_idx in (1, 2, 3, 6, 10): cell.alignment = align_center
        else: cell.alignment = align_left
        if c_idx == 10 and val == 'PASSED':
            cell.font = font_pass
            cell.fill = fill_pass

ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 14
ws2.column_dimensions['D'].width = 28
ws2.column_dimensions['E'].width = 45
ws2.column_dimensions['F'].width = 14
ws2.column_dimensions['G'].width = 45
ws2.column_dimensions['H'].width = 45
ws2.column_dimensions['I'].width = 35
ws2.column_dimensions['J'].width = 14
ws2.column_dimensions['K'].width = 12

# SHEET 3: Khuyến Nghị & Kiểm Soát Lỗi
ws3 = wb.create_sheet(title='Khuyến Nghị & Kiểm Soát Lỗi')
ws3.views.sheetView[0].showGridLines = True

ws3.cell(2, 2, 'DANH MỤC KHUYẾN NGHỊ & KIỂM SOÁT NGOẠI LỆ SPRINT 2').font = font_title
ws3.cell(3, 2, 'Tổng hợp các giải pháp xử lý ngoại lệ nghiệp vụ theo quy tắc GEMINI.md').font = font_subtitle

headers_s3 = ['STT', 'Hạng mục kiểm soát', 'Nguy cơ tiềm ẩn / Lỗi thực tế', 'Giải pháp đã triển khai & Khuyến nghị', 'Mức độ', 'Trạng thái']
for c_idx, h in enumerate(headers_s3, start=2):
    cell = ws3.cell(5, c_idx, h)
    cell.font = font_header
    cell.fill = fill_blue_header
    cell.alignment = align_center
    cell.border = header_border

rec_data = [
    (1, 'Kiểm tra ràng buộc thời gian', 'Ngày kết thúc chương trình thực tập nhỏ hơn ngày bắt đầu làm sai lệch kế hoạch.', 'Bắt buộc áp dụng validation decorator: ngay_ket_thuc > ngay_bat_dau, trả HTTP 422.', 'Cao', 'Đã kiểm thử ĐẠT'),
    (2, 'Chống trùng lặp email TTS', 'Cập nhật trùng email gây xung đột tài khoản đăng nhập người dùng.', 'Truy vấn kiểm tra email tồn tại trong DB trước khi UPDATE, ném HTTP 400 Bad Request.', 'Nghiêm trọng', 'Đã kiểm thử ĐẠT'),
    (3, 'Kiểm soát khóa ngoại (Foreign Key)', 'Tạo chương trình với phòng ban không tồn tại hoặc nộp báo cáo với mã hồ sơ ảo.', 'Kiểm tra sự tồn tại của record quan hệ trước khi commit, trả HTTP 400/404 rõ ràng.', 'Cao', 'Đã kiểm thử ĐẠT'),
    (4, 'Giới hạn độ dài nhận xét / email', 'Nhập ghi chú từ chối hoặc nhận xét quá dài làm vỡ định dạng email / database buffer.', 'Giới hạn max_length <= 500 ký tự trong schema Pydantic, báo lỗi HTTP 422.', 'Trung bình', 'Đã kiểm thử ĐẠT'),
    (5, 'Độc lập môi trường kiểm thử', 'Chạy test phụ thuộc vào port 8000 và database MySQL môi trường production.', 'Refactor sang FastAPI TestClient với SQLite Isolated Engine và mock request session.', 'Nghiêm trọng', 'Đã xử lý 100%'),
]

for r_idx, row_vals in enumerate(rec_data, start=6):
    fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    for c_idx, val in enumerate(row_vals, start=2):
        cell = ws3.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        if fill.fill_type: cell.fill = fill
        if c_idx in (2, 6, 7): cell.alignment = align_center
        else: cell.alignment = align_left
        if c_idx == 7:
            cell.font = font_pass
            cell.fill = fill_pass

ws3.column_dimensions['A'].width = 3
ws3.column_dimensions['B'].width = 6
ws3.column_dimensions['C'].width = 30
ws3.column_dimensions['D'].width = 45
ws3.column_dimensions['E'].width = 50
ws3.column_dimensions['F'].width = 16
ws3.column_dimensions['G'].width = 20

out_path = 'Bao_Cao_Kiem_Thu_Sprint2_Tran_Quoc_Huy.xlsx'
wb.save(out_path)
print(f'Successfully generated {out_path}')
